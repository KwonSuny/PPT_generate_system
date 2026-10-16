import copy
import json
import os
import tempfile
import unittest
import zipfile
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from pptx import Presentation
from pptx.oxml.ns import qn
from pydantic import ValidationError

import main
from agents.writer import get_writer_agent
from agents.planner import get_planner_agent
from agents.template import get_layout_agent, select_layout_canvases, composition_guidance
from config.content import (ContentConstraintError, bind_canvas, cell_text, identified_slide,
                            lock_slide, lock_text_canvas, parse_page_policy)
from config.schema import PPTFinalBlueprint, PPTOutline, SlideCanvas, SlideElement, TableSpec
from config.sanitizers import normalize_outline_payload, normalize_blueprint_payload, parse_completion_payload
from config.state import PPTGenerationState
from renderer import PPTXRenderer
from test_pipeline import outline_fixture, blueprint_fixture


def table_slide():
    slide = outline_fixture()["slides"][0]
    slide.update(title="指标对比", layout_type="table", bullet_points=[], full_content=[], table={
        "columns": [{"id": "metric", "title": "指标"}, {"id": "before", "title": "基准值"},
                    {"id": "after", "title": "调整值"}],
        "rows": [{"id": "accuracy", "cells": [{"value": "准确率"}, {"value": 0.8765, "format": {"style": "percent", "decimals": 2}},
                                                {"value": 0.9132, "format": {"style": "percent", "decimals": 2}}]},
                 {"id": "delta", "cells": [{"value": "差值"}, {"value": -12.345, "format": {"style": "decimal", "decimals": 2}},
                                             {"value": "-0.03150"}]}],
        "source": "合成回归数据，不代表真实实验结果",
    })
    return identified_slide(slide)


def diagram_slide(direction="LR", branch=False):
    slide = outline_fixture()["slides"][0]
    diagram = {"direction": direction,
               "nodes": [{"id": "input", "label": "用户需求"}, {"id": "plan", "label": "内容规划"}, {"id": "review", "label": "页面审核"}],
               "edges": [{"id": "first", "source": "input", "target": "plan", "label": "生成"},
                         {"id": "second", "source": "plan", "target": "review", "label": "检验"}]}
    if branch:
        diagram["nodes"].append({"id": "alternative", "label": "另一分支"})
        diagram["edges"].append({"id": "branch", "source": "input", "target": "alternative", "label": "备选"})
    slide.update(title="生产流程", layout_type="diagram", bullet_points=[], full_content=[], diagram=diagram)
    return identified_slide(slide)


def canvas_for(slide):
    elements = [{"type": "textbox", "left": 0.7, "top": 0.45, "width": 11.8, "height": 0.9,
                 "paragraphs": [{"content_ref": "title", "text": "由程序回填", "font_size": 32, "bold": True}]}]
    if slide.get("layout_type") in {"table", "diagram", "svg"}:
        elements.append({"type": slide["layout_type"], "left": 0.7, "top": 1.65, "width": 11.8, "height": 4.9})
    else:
        for i, text in enumerate(slide.get("full_content", [])):
            elements.append({"type": "textbox", "left": 0.7, "top": 1.7 + i, "width": 10, "height": 0.8,
                             "paragraphs": [{"content_ref": f"full_content/point-{i + 1}", "text": text, "font_size": 18}]})
    return bind_canvas({"elements": elements}, slide)


def response(payload):
    return SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content=json.dumps(payload, ensure_ascii=False)))])


class ContentContractTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix="ppt_content_test_")
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.output = self.root / "result.pptx"
        replacement = patch.dict(os.environ, {"DASHSCOPE_API_KEY": "test-key-not-a-credential"})
        replacement.start()
        self.addCleanup(replacement.stop)

    def render(self, canvas, name="result.pptx"):
        path = self.root / name
        main._render_blueprint(PPTFinalBlueprint(schema_version=2, canvases=[canvas]), str(path))
        return Presentation(path)

    def _tight_layout_case(self):
        slide = identified_slide({**outline_fixture()["slides"][0], "title": "种植与采摘", "punchline": "",
                                  "bullet_points": ["主要产区位于赤道南北回归线之间（咖啡带）"],
                                  "full_content": ["全球咖啡种植集中在热带区域，这一地带被称为咖啡带。"]})
        canvas = canvas_for(slide)
        canvas.elements.append(canvas.elements[1].model_copy(deep=True))
        bullet, full = canvas.elements[1:]
        bullet.id = "point-1-bullet"; bullet.top = 2.3; bullet.height = 0.4
        bullet.paragraphs[0].content_ref = "bullet_points/point-1"; bullet.paragraphs[0].font_size = 20
        full.id = "point-1-full"; full.top = 2.7; full.height = 0.6
        canvas = bind_canvas(canvas, slide)
        return slide, canvas

    def test_capacity_error_returns_feedback_and_preserves_text_and_font(self):
        slide, bad = self._tight_layout_case()
        state = PPTGenerationState(topic="咖啡")
        calls = []
        def designer(state, payload):
            calls.append(copy.deepcopy(payload))
            canvas = bad.model_copy(deep=True)
            if len(calls) > 1:
                canvas.elements[1].height = 0.52
                canvas.elements[2].top = 2.87
            return PPTFinalBlueprint(canvases=[canvas])
        result = main._layout_page(state, designer, {"slides": [slide]}, 2, str(self.root))
        self.assertEqual(len(calls), 2)
        feedback = calls[1]["layout_feedback"]
        self.assertIn("point-1-bullet", feedback["error"])
        self.assertGreater(feedback["required_text_heights"][1]["minimum_height"], 0.49)
        self.assertEqual(result.elements[1].paragraphs[0].font_size, 20)
        self.assertEqual(result.elements[2].paragraphs[0].text, slide["full_content"][0])
        self.assertTrue((self.root / "layout_page_2_attempt_1.json").is_file())
        self.assertTrue(result.layout_resolved)

    def test_capacity_repair_is_bounded_and_rejects_font_shrinking(self):
        slide, bad = self._tight_layout_case()
        state = PPTGenerationState(topic="咖啡",requirements='保持字号不变')
        stages = []
        def executor(key, inputs, operation, *args):
            stages.append(key)
            return operation()
        state.stage_executor = executor
        def designer(state, payload):
            candidate = bad.model_copy(deep=True)
            if "layout_feedback" in payload:
                candidate.elements[1].paragraphs[0].font_size = 18
                candidate.elements[1].height = 0.55
                candidate.elements[2].top = 2.9
            return PPTFinalBlueprint(canvases=[candidate])
        result = main._layout_page(state, designer, {"slides": [slide]}, 2, str(self.root))
        self.assertEqual(stages, ["layout/page2", "layout/page2/repair1", "layout/page2/repair2", "layout/page2/fallback"])
        sizes = {p.content_ref:p.font_size for e in result.elements for p in e.paragraphs or []}
        self.assertGreaterEqual(sizes['bullet_points/point-1'], 20)
        self.assertEqual(state.get('layout_fallbacks')[0]['page'], 2)
        with self.assertRaisesRegex(ContentConstraintError, "缩小.*字号"):
            main._layout_page(PPTGenerationState(topic="咖啡"), designer, {"slides": [slide]}, 2, str(self.root), stage_prefix="layout/candidate-clear")

    def test_layout_repair_cannot_drop_previously_displayed_bullet(self):
        slide, bad = self._tight_layout_case()
        calls = []
        def designer(state, payload):
            calls.append(payload)
            candidate = bad.model_copy(deep=True)
            if "layout_feedback" in payload:
                candidate.elements.pop(1)
            return PPTFinalBlueprint(canvases=[candidate])
        result = main._layout_page(PPTGenerationState(topic="咖啡"), designer, {"slides": [slide]}, 2, str(self.root))
        self.assertEqual(len(calls), 3)
        self.assertEqual({p.content_ref:p.text for e in result.elements for p in e.paragraphs or []}['bullet_points/point-1'],slide['bullet_points'][0])

    def test_stable_layout_keeps_native_table_diagram_and_svg_and_rejects_impossible_text(self):
        from config.layout_fallback import stable_layout
        from config.visual import rule_issues
        for slide in (table_slide(), diagram_slide()):
            with self.subTest(kind=slide['layout_type']):
                canvas, _ = stable_layout(slide)
                self.assertEqual(rule_issues(canvas), [])
                native = next(e for e in canvas.elements if e.type == slide['layout_type'])
                self.assertEqual(getattr(native,native.type).model_dump(),slide[native.type])
                self.render(canvas)
        slide = copy.deepcopy(outline_fixture()['slides'][0])
        slide['svg_code'] = '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 800 400"><text x="20" y="60">必须保留的图示</text></svg>'
        canvas, _ = stable_layout(slide)
        self.assertEqual(next(e.svg_code for e in canvas.elements if e.type=='svg'),slide['svg_code'])
        slide['full_content'] = ['不可丢失的事实与数字' * 2000]
        slide['bullet_points'] = ['完整内容']
        with self.assertRaisesRegex(ContentConstraintError,'未删改原文'):
            stable_layout(slide)
        huge_table = table_slide()
        huge_table['table']['rows'][0]['cells'][0]['value'] = '必须保留的表格数据' * 2000
        with self.assertRaises(ContentConstraintError) as caught:
            stable_layout(huge_table)
        self.assertTrue(caught.exception.capacity_conflict)
        cyclic = diagram_slide()
        cyclic['diagram']['edges'].append({'id':'loop','source':'review','target':'input'})
        with self.assertRaises(ContentConstraintError) as caught:
            stable_layout(cyclic)
        self.assertFalse(caught.exception.capacity_conflict)

    def test_fallback_stage_is_local_and_successful_revision_clears_page_warning(self):
        slide, bad = self._tight_layout_case()
        state = PPTGenerationState(topic="主题")
        stages = []
        def executor(key, inputs, operation, model, result_model, external):
            stages.append((key,external))
            return operation()
        state.stage_executor = executor
        main._layout_page(state,lambda *_: PPTFinalBlueprint(canvases=[bad]),{'slides':[slide]},2,str(self.root))
        self.assertEqual(stages[-1],('layout/page2/fallback',False))
        good = bad.model_copy(deep=True)
        good.elements[1].height = .9; good.elements[2].top = 3.3; good.elements[2].height = .9
        main._layout_page(state,lambda *_: PPTFinalBlueprint(canvases=[good]),{'slides':[slide]},2,str(self.root))
        self.assertEqual(state.get('layout_fallbacks'),[])

    def test_layout_unknown_failure_is_not_repaired(self):
        state = PPTGenerationState(topic="咖啡")
        from config.jobs import StageOutcomeUnknown
        def designer(*args):
            raise StageOutcomeUnknown("网络结果未知")
        with patch.object(main, "get_layout_agent", return_value=designer):
            with self.assertRaises(StageOutcomeUnknown):
                main._layout_page(state, designer, {"slides": [table_slide()]}, 1, str(self.root))
        self.assertFalse(state.progress_log)

    def test_writer_svg_can_coexist_with_native_table_and_survive_binding(self):
        slide = table_slide()
        source = '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 200 100"><rect width="200" height="100" fill="#123456"/></svg>'
        output = copy.deepcopy(slide); output["svg_code"] = source
        with patch("agents.writer.create_client"), patch("agents.writer.chat_completion_with_retry", return_value=response({"slides": [output]})):
            result = get_writer_agent("test-key-not-a-credential")(PPTGenerationState(topic="主题", requirements="用表格和内联SVG表达"), {"slides": [slide]})
        self.assertEqual(result.slides[0].svg_code, source)
        self.assertEqual(result.slides[0].table, SlideElement.model_validate(slide).table)
        canvas = canvas_for(slide).model_dump()
        canvas["elements"].append({"type": "svg", "left": 0.7, "top": 6.7, "width": 1, "height": 0.5})
        bound = bind_canvas(canvas, result.slides[0])
        self.assertEqual(bound.elements[-1].svg_code, source)
        self.assertEqual(bound.elements[1].type, "table")

    def _crowded_columns(self):
        elements = [{"id": "diagram", "type": "svg", "left": 0.8, "top": 1.8, "width": 11.7, "height": 1.6,
                     "svg_code": '<svg viewBox="0 0 100 20"/>'}]
        for index, (left, top) in enumerate([(1.2, 3.7), (7.2, 3.7), (1.2, 5.7), (7.2, 5.7)]):
            elements.append({"id": f"title-{index}", "type": "textbox", "left": left, "top": top, "width": 4.8,
                             "height": 0.9 if index == 3 else 0.55,
                             "paragraphs": [{"text": "保持对应要点标题的文字完整" + ("和清楚可读" if index == 3 else ""), "font_size": 20}]})
            elements.append({"id": f"body-{index}", "type": "textbox", "left": left, "top": 6.65 if index == 3 else top + 0.6,
                             "width": 4.8, "height": 1.15,
                             "paragraphs": [{"text": "这里使用固定的解释文字验证容量，保持内容和数字不变，不能靠缩小字号掩盖越界问题。", "font_size": 18}]})
        return SlideCanvas.model_validate({"elements": elements})

    def test_final_candidate_reflow_preserves_columns_fonts_order_and_art(self):
        raw = self._crowded_columns()
        with self.assertRaises(ContentConstraintError):
            PPTXRenderer.resolve_canvas(raw, fit_text=True)
        resolved = PPTXRenderer.resolve_canvas(raw, fit_text=True, reflow_text=True)
        self.assertEqual(resolved.elements[0], raw.elements[0])
        for before, after in zip(raw.elements, resolved.elements):
            self.assertEqual((before.left, before.width), (after.left, after.width))
            if before.type == "textbox":
                self.assertEqual([(p.text, p.font_size) for p in before.paragraphs], [(p.text, p.font_size) for p in after.paragraphs])
                self.assertGreaterEqual(after.top, 3.55)
                self.assertLessEqual(after.top + after.height, 7.35)
                self.assertLessEqual(abs(after.top - before.top), 1)
        self.assertEqual(resolved.elements[1].top, resolved.elements[3].top)
        self.assertEqual(resolved.elements[5].top, resolved.elements[7].top)
        self.assertEqual(raw.elements[-1].top, 6.65)
        raw.elements[2].paragraphs[0].text = "短说明"
        varied = PPTXRenderer.resolve_canvas(raw, fit_text=True, reflow_text=True)
        self.assertEqual(varied.elements[5].top, varied.elements[7].top)

    def test_final_candidate_reflow_rejects_impossible_capacity_and_horizontal_bounds(self):
        raw = self._crowded_columns()
        raw.elements[-1].paragraphs[0].text = "容量必须可读" * 300
        with self.assertRaises(ContentConstraintError):
            PPTXRenderer.resolve_canvas(raw, fit_text=True, reflow_text=True)
        raw = self._crowded_columns(); raw.elements[-1].left = 12
        with self.assertRaises(ContentConstraintError):
            PPTXRenderer.resolve_canvas(raw, fit_text=True, reflow_text=True)

    def test_composition_guidance_uses_semantics_and_preserves_input(self):
        slide = diagram_slide()
        slide["full_content"] = ["简短说明"] * 4
        original = copy.deepcopy(slide)
        guidance = composition_guidance(slide)
        self.assertEqual(guidance["candidates"][0], "上方关系图、下方二列解释")
        self.assertEqual(slide, original)
        self.assertFalse(guidance["art_is_separate"])
        cover = {"layout_type": "title_page", "svg_code": "<svg/>"}
        self.assertEqual(composition_guidance(cover)["candidates"][0], "左侧标题与结论、右侧独立主题图")
        table = composition_guidance(table_slide())
        self.assertEqual((table["table_rows"], table["table_columns"]), (2, 3))

    def test_layout_matches_explicit_ids_without_expanding_single_page_request(self):
        slide = table_slide()
        requested = {"slide_id": slide["id"], "elements": []}
        extra = {"slide_id": "other-page", "elements": []}
        state = PPTGenerationState(topic="单页排版")
        payload = {"schema_version": 2, "canvases": [extra, requested]}
        result = select_layout_canvases(payload, [slide], state)
        self.assertEqual(result["canvases"], [requested])
        self.assertEqual(len(payload["canvases"]), 2)
        self.assertIn("忽略 1 张", state.progress_log[-1])
        for canvases in ([requested, requested], [extra], [{"elements": []}, requested], [{"elements": []}, {"elements": []}]):
            with self.subTest(canvases=canvases), self.assertRaises(ContentConstraintError):
                select_layout_canvases({"canvases": canvases}, [slide])
        self.assertEqual(select_layout_canvases({"canvases": [{"elements": []}]}, [slide])["canvases"], [{"elements": []}])

    def test_page_policy_defaults_exact_range_chinese_and_ordinals(self):
        for raw, text, expected in [("", "普通汇报", (2, 5)), ("1", "", (1, 1)), ("6–8", "", (6, 8)),
                                    ("", "生成三页", (3, 3)), ("", "控制在 2 至 4 页", (2, 4)),
                                    ("", "页数：6", (6, 6)), ("", "第 2–3 页参考旧版", (2, 5)),
                                    ("4", "制作八页", (4, 4))]:
            with self.subTest(raw=raw, text=text):
                policy = parse_page_policy(raw, text)
                self.assertEqual((policy.min_slides, policy.max_slides), expected)
        for raw in ("0", "51", "5-2", "1.5", "三十十", "1," , "9" * 100):
            with self.subTest(raw=raw), self.assertRaises(ContentConstraintError):
                parse_page_policy(raw)
        with self.assertRaises(ContentConstraintError):
            parse_page_policy(text="生成 3 页或 8 页")

    def test_offline_demo_honors_explicit_count_and_default(self):
        for raw, count in [("", 3), ("1", 1), ("2", 2), ("6", 6), ("4–6", 4)]:
            state = PPTGenerationState(topic="页数测试")
            with self.subTest(raw=raw), patch.dict(os.environ, {"DASHSCOPE_API_KEY": ""}):
                main.run_ppt_system("页数测试", filename=str(self.output), state=state, page_range=raw)
                self.assertEqual(len(Presentation(self.output).slides), count)
                self.assertEqual(state.get("result_kind"), "offline_demo")

    def test_out_of_range_output_stops_before_writer_images_and_export(self):
        state = PPTGenerationState(topic="需要六页")
        outline = PPTOutline.model_validate({"topic": "主题", "slides": outline_fixture()["slides"][:1]})
        main._render_offline_demo("已有文件", str(self.output))
        original = self.output.read_bytes()
        with patch.object(main, "get_planner_agent", return_value=lambda *a, **k: outline), patch.object(main, "get_writer_agent") as writer, patch.object(main, "get_image_generator") as image:
            self.assertIsNone(main.run_ppt_system("主题", filename=str(self.output), state=state, page_range="6", run_dir=str(self.root)))
        writer.return_value.assert_not_called()
        image.return_value.assert_not_called()
        self.assertEqual(self.output.read_bytes(), original)
        self.assertTrue(any("实际生成 1 页" in message for message in state.progress_log))

    def test_full_pipeline_accepts_more_than_five_and_persists_resolved_layout(self):
        slides = outline_fixture()["slides"] * 2
        outline = PPTOutline.model_validate({"topic": "六页主题", "slides": slides})
        state = PPTGenerationState(topic="六页主题")
        def layout(state, payload):
            return PPTFinalBlueprint(canvases=[canvas_for(payload["slides"][0])])
        with patch.object(main, "get_planner_agent", return_value=lambda *a, **k: outline), patch.object(main, "get_writer_agent", return_value=lambda *a, **k: outline), patch.object(main, "get_layout_agent", return_value=layout), patch.object(main, "get_image_generator"):
            result = main.run_ppt_system("主题", filename=str(self.output), state=state, page_range="6", run_dir=str(self.root))
        self.assertEqual(result, str(self.output))
        self.assertEqual(len(Presentation(self.output).slides), 6)
        self.assertEqual(state.get("outline")["schema_version"], 2)
        self.assertEqual(len({s["id"] for s in state.get("outline")["slides"]}), 6)
        self.assertTrue(all(c["layout_resolved"] for c in state.get("blueprint")["canvases"]))

    def test_native_table_keeps_values_units_precision_and_editable_cells(self):
        slide = table_slide()
        canvas = PPTXRenderer.resolve_canvas(canvas_for(slide))
        presentation = self.render(canvas)
        table = next(shape.table for shape in presentation.slides[0].shapes if shape.has_table)
        self.assertEqual([[cell.text for cell in row.cells] for row in table.rows],
                         [["指标", "基准值", "调整值"], ["准确率", "87.65%", "91.32%"], ["差值", "-12.35", "-0.03150"]])
        self.assertEqual(canvas.elements[1].table.rows[0].cells[1].value, 0.8765)
        self.assertEqual(table.cell(1, 1)._tc.get_or_add_tcPr().get("anchor"), "ctr")
        self.assertEqual(table.cell(1, 1).margin_top, int(0.08 * 914400))
        table.cell(1, 1).text = "可编辑"
        edited = self.root / "edited.pptx"
        presentation.save(edited)
        reopened = Presentation(edited)
        self.assertEqual(next(shape.table for shape in reopened.slides[0].shapes if shape.has_table).cell(1, 1).text, "可编辑")

    def test_table_data_rejects_missing_cells_duplicate_ids_nonfinite_and_boolean(self):
        base = table_slide()["table"]
        variants = []
        missing = copy.deepcopy(base); missing["rows"][0]["cells"].pop(); variants.append(missing)
        duplicate = copy.deepcopy(base); duplicate["columns"][1]["id"] = duplicate["columns"][0]["id"]; variants.append(duplicate)
        for value in (float("nan"), float("inf"), True, {"unexpected": 1}):
            invalid = copy.deepcopy(base); invalid["rows"][0]["cells"][0]["value"] = value; variants.append(invalid)
        for value in variants:
            with self.subTest(value=value), self.assertRaises(ValidationError):
                TableSpec.model_validate(value)

    def test_table_overflow_fails_without_losing_rows_or_shrinking(self):
        slide = table_slide()
        slide["table"]["rows"][0]["cells"][0]["value"] = "中文长文本" * 200
        canvas = canvas_for(slide)
        original = canvas.model_dump()
        with self.assertRaisesRegex(ContentConstraintError, "表格文字需要"):
            PPTXRenderer.resolve_canvas(canvas)
        self.assertEqual(canvas.model_dump(), original)

    def test_connected_arrows_have_native_node_ports_and_no_detached_triangle(self):
        for direction in ("LR", "TB"):
            with self.subTest(direction=direction):
                canvas = PPTXRenderer.resolve_canvas(canvas_for(diagram_slide(direction)))
                shapes = self.render(canvas).slides[0].shapes
                nodes = {shape.name.rsplit("-node-", 1)[-1]: shape for shape in shapes if "-node-" in shape.name}
                connectors = [shape for shape in shapes if "-edge-" in shape.name]
                self.assertEqual(len(connectors), 2)
                for connector, source, target in zip(connectors, ("input", "plan"), ("plan", "review")):
                    properties = connector._element.nvCxnSpPr.cNvCxnSpPr
                    start = properties.find(qn("a:stCxn")); end = properties.find(qn("a:endCxn"))
                    self.assertEqual(int(start.get("id")), nodes[source].shape_id)
                    self.assertEqual(int(end.get("id")), nodes[target].shape_id)
                    self.assertEqual((int(start.get("idx")), int(end.get("idx"))), (3, 1) if direction == "LR" else (2, 0))
                    self.assertEqual(connector._element.spPr.ln.find(qn("a:tailEnd")).get("type"), "triangle")
                    self.assertLessEqual(abs(connector.begin_x - (nodes[source].left + nodes[source].width if direction == "LR" else nodes[source].left + nodes[source].width / 2)), 1)
                    self.assertLessEqual(abs(connector.end_y - (nodes[target].top + nodes[target].height / 2 if direction == "LR" else nodes[target].top)), 1)
                self.assertEqual(len([shape for shape in shapes if shape._element.tag == qn("p:sp") and "-node-" in shape.name]), 3)

    def test_branch_labels_do_not_overlap_and_cycle_or_cross_layer_is_explicit(self):
        for direction in ("LR", "TB"):
            canvas = PPTXRenderer.resolve_canvas(canvas_for(diagram_slide(direction, branch=True)))
            labels = [e.label_box for e in canvas.elements[1].diagram_geometry.edges if e.label_box]
            self.assertFalse(any(PPTXRenderer._boxes_overlap(a, b) for i, a in enumerate(labels) for b in labels[i + 1:]))
        slide = diagram_slide()
        slide["diagram"]["edges"].append({"id": "cycle", "source": "review", "target": "input"})
        with self.assertRaisesRegex(ContentConstraintError, "环路"):
            PPTXRenderer.resolve_canvas(canvas_for(slide))
        slide = diagram_slide()
        slide["diagram"]["edges"].append({"id": "skip", "source": "input", "target": "review"})
        with self.assertRaisesRegex(ContentConstraintError, "跨层"):
            PPTXRenderer.resolve_canvas(canvas_for(slide))

    def test_svg_retains_code_and_non_square_aspect_ratio(self):
        slide = identified_slide({**outline_fixture()["slides"][0], "title": "比例检查", "layout_type": "svg", "bullet_points": [], "full_content": [],
                                  "svg_code": '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 300 100"><rect width="300" height="100" fill="#005088"/></svg>'})
        canvas = PPTXRenderer.resolve_canvas(canvas_for(slide))
        shapes = self.render(canvas).slides[0].shapes
        picture = next(shape for shape in shapes if shape.shape_type == 13)
        self.assertAlmostEqual(picture.width / picture.height, 3, places=5)
        self.assertEqual(canvas.elements[1].svg_code, slide["svg_code"])

    def test_invalid_svg_or_geometry_cannot_export_as_success(self):
        slide = identified_slide({**outline_fixture()["slides"][0], "layout_type": "svg", "bullet_points": [], "full_content": [], "svg_code": '<svg viewBox="0 0 200 100"><broken></svg>'})
        canvas = PPTXRenderer.resolve_canvas(canvas_for(slide))
        with self.assertRaisesRegex(ContentConstraintError, "SVG 渲染失败"):
            self.render(canvas)
        self.assertFalse(self.output.exists())
        bad = canvas_for(table_slide()); bad.elements[1].left = -1
        with self.assertRaisesRegex(ContentConstraintError, "超出画布"):
            PPTXRenderer.resolve_canvas(bad)

    def test_fresh_text_can_expand_into_free_space_but_saved_text_cannot(self):
        slide = table_slide()
        canvas = canvas_for(slide)
        canvas.elements[0].paragraphs[0].font_size = 40
        canvas.elements[0].height = 0.8
        with self.assertRaisesRegex(ContentConstraintError, "超出容量"):
            PPTXRenderer.resolve_canvas(canvas)
        resolved = PPTXRenderer.resolve_canvas(canvas, fit_text=True)
        self.assertGreater(resolved.elements[0].height, 0.8)
        self.assertEqual(resolved.elements[0].paragraphs[0].font_size, 40)
        self.assertEqual(resolved.elements[1].top, canvas.elements[1].top)
        with self.assertRaisesRegex(ContentConstraintError, "超出容量"):
            lock_text_canvas(resolved, slide, {**slide, "title": "中文长标题" * 30})
        canvas.elements[1].top = 1.27
        with self.assertRaisesRegex(ContentConstraintError, "超出容量"):
            PPTXRenderer.resolve_canvas(canvas, fit_text=True)

    def test_text_lock_skips_images_layout_and_preserves_non_target_json_and_slide_xml(self):
        state = PPTGenerationState(topic="固定内容")
        outline, blueprint = outline_fixture(), blueprint_fixture()
        for canvas in blueprint["canvases"]:
            for element in canvas["elements"]:
                for paragraph in element["paragraphs"]:
                    paragraph.pop("content_ref")
            canvas["elements"].append({"type": "icon", "left": 0.4, "top": 3, "width": 0.4, "height": 0.4, "icon_name": "circle"})
        state.set("outline", copy.deepcopy(outline)); state.set("blueprint", copy.deepcopy(blueprint)); state.set_operation("incremental_page", "text", pages=[2])
        original_path = self.root / "before.pptx"
        main._render_blueprint(PPTFinalBlueprint.model_validate(blueprint), str(original_path))
        def writer(state, payload):
            changed = copy.deepcopy(payload)
            changed["slides"][0].update(title="调整后的标题", full_content=["调整后的正文"], image_prompt="不应生成图", theme={"primary": "#FF0000", "secondary": "#FF0000", "content": "#FF0000", "background": "#000000"})
            return PPTOutline.model_validate(changed)
        with patch.object(main, "get_writer_agent", return_value=writer), patch.object(main, "get_layout_agent") as layout, patch.object(main, "get_image_generator") as image:
            self.assertEqual(main.run_incremental_modify_pages("修改文字", [2], filename=str(self.output), state=state, run_dir=str(self.root)), str(self.output))
        layout.assert_not_called(); image.assert_not_called()
        self.assertEqual(state.get("outline")["slides"][1]["theme"], outline["slides"][1]["theme"])
        for index in (0, 2):
            self.assertEqual(state.get("outline")["slides"][index], outline["slides"][index])
            self.assertEqual(state.get("blueprint")["canvases"][index], blueprint["canvases"][index])
        with zipfile.ZipFile(original_path) as before, zipfile.ZipFile(self.output) as after:
            for number in (1, 3):
                self.assertEqual(before.read(f"ppt/slides/slide{number}.xml"), after.read(f"ppt/slides/slide{number}.xml"))
        before, after = Presentation(original_path).slides[1], Presentation(self.output).slides[1]
        self.assertEqual([(s.left, s.top, s.width, s.height) for s in before.shapes], [(s.left, s.top, s.width, s.height) for s in after.shapes])

    def test_text_lock_capacity_and_ambiguous_legacy_mapping_fail_atomically(self):
        for ambiguous in (False, True):
            state = PPTGenerationState(topic="固定内容")
            outline, blueprint = outline_fixture(), blueprint_fixture()
            if ambiguous:
                outline["slides"][1]["title"] = outline["slides"][1]["full_content"][0] = "重复内容"
                for element in blueprint["canvases"][1]["elements"]:
                    element["paragraphs"][0] = {"text": "重复内容"}
            state.set("outline", copy.deepcopy(outline)); state.set("blueprint", copy.deepcopy(blueprint)); state.set_operation("incremental_page", "text", pages=[2])
            def writer(state, payload):
                changed = copy.deepcopy(payload)
                changed["slides"][0]["title"] = "新标题" if ambiguous else "中文长文本" * 200
                return PPTOutline.model_validate(changed)
            with self.subTest(ambiguous=ambiguous), patch.object(main, "get_writer_agent", return_value=writer), patch.object(main, "get_layout_agent") as layout:
                self.assertIsNone(main.run_incremental_modify_pages("修改", [2], filename=str(self.output), state=state, run_dir=str(self.root)))
                layout.assert_not_called()
                self.assertEqual(state.get("outline"), outline); self.assertEqual(state.get("blueprint"), blueprint)
                self.assertTrue(any(("无法唯一映射" if ambiguous else "超出容量") in p for p in state.progress_log))
        self.assertFalse(self.output.exists())

    def test_table_and_diagram_text_updates_keep_physical_geometry_and_styles(self):
        for old in (table_slide(), diagram_slide()):
            canvas = PPTXRenderer.resolve_canvas(canvas_for(old))
            changed = copy.deepcopy(old)
            if old["table"]:
                changed["table"]["rows"][0]["cells"][1]["value"] = 0.9325
                changed["table"]["source"] = "合成数据更新"
            else:
                changed["diagram"]["nodes"][0]["label"] = "最新需求"
                changed["diagram"]["edges"][0]["label"] = ""
            locked = lock_text_canvas(canvas, old, lock_slide(old, changed, "text"))
            old_element, new_element = canvas.elements[1], locked.elements[1]
            for name in ("left", "top", "width", "height", "table_style", "column_widths", "row_heights", "diagram_style", "diagram_geometry"):
                self.assertEqual(getattr(old_element, name), getattr(new_element, name))
            before = self.render(canvas, "before.pptx").slides[0].shapes
            after = self.render(locked).slides[0].shapes
            self.assertEqual([(s.name, s.left, s.top, s.width, s.height) for s in before], [(s.name, s.left, s.top, s.width, s.height) for s in after])

    def test_text_cannot_change_structure_and_design_cannot_change_values(self):
        old = table_slide(); changed = copy.deepcopy(old)
        changed["table"]["rows"].append(copy.deepcopy(changed["table"]["rows"][0]))
        with self.assertRaisesRegex(ContentConstraintError, "增删"):
            lock_slide(old, changed, "text")
        changed = copy.deepcopy(old); changed["table"]["rows"][0]["cells"][1]["value"] = 99
        locked = lock_slide(old, changed, "design")
        self.assertEqual(locked["table"], old["table"])
        old = diagram_slide(); changed = copy.deepcopy(old)
        changed["diagram"]["nodes"][0]["label"] = "不应改写"
        self.assertEqual(lock_slide(old, changed, "design")["diagram"], old["diagram"])
        old["svg_code"] = '<svg viewBox="0 0 100 100"><text>数据 12.5</text></svg>'
        changed["svg_code"] = '<svg viewBox="0 0 100 100"><text>数据 99</text></svg>'
        self.assertEqual(lock_slide(old, changed, "design")["svg_code"], old["svg_code"])
        old = outline_fixture()["slides"][0]
        old["full_content"].append("旧版中额外保留的正文")
        result = normalize_outline_payload({"slides": [{}]}, skeleton={"slides": [old]}, lock_scope="design")
        self.assertEqual(result["slides"][0]["full_content"], old["full_content"])

    def test_writer_structure_keeps_existing_native_evidence(self):
        for old in (table_slide(), diagram_slide()):
            changed = copy.deepcopy(old)
            if old["table"]:
                changed["table"]["rows"][0]["cells"][1]["value"] = 99
            else:
                changed["diagram"]["nodes"][0]["label"] = "模型改写"
            result = normalize_outline_payload({"slides": [changed]}, skeleton={"slides": [old]}, lock_scope="structure")
            self.assertEqual(result["slides"][0]["table"], old["table"])
            self.assertEqual(result["slides"][0]["diagram"], old["diagram"])

    def test_planner_leaves_native_data_to_writer_and_retains_raw_context(self):
        slide = table_slide()
        slide["table"] = {"headers": ["方案", "准确率"], "rows": [["A", "87.65%"]]}
        payload = {"topic": "主题", "slides": [slide]}
        state = PPTGenerationState(topic="主题", requirements="方案 A 的准确率为 87.65%")
        with patch("agents.planner.create_client"), patch("agents.planner.chat_completion_with_retry", return_value=response(payload)):
            outline = get_planner_agent("test-key-not-a-credential")(state)
        self.assertIsNone(outline.slides[0].table)
        self.assertEqual(outline.slides[0].layout_type, "table")
        self.assertIn("87.65%", state.messages[-1]["content"])

    def test_native_preview_uses_semantic_title_and_includes_table_cells(self):
        from app import _build_ppt_preview
        slide = table_slide()
        canvas = PPTXRenderer.resolve_canvas(canvas_for(slide))
        canvas.elements = list(reversed(canvas.elements))
        self.render(canvas)
        preview = _build_ppt_preview(str(self.output), {"slides": [slide]})
        self.assertEqual(preview[0]["title"], slide["title"])
        self.assertIn("87.65%", preview[0]["text"])
        self.assertIn("差值", preview[0]["text"])

    def test_writer_rejects_extra_pages_instead_of_silent_truncation_or_fallback(self):
        skeleton = {"topic": "主题", "slides": [table_slide()]}
        raw = {"topic": "主题", "slides": skeleton["slides"] * 2}
        with patch("agents.writer.create_client"), patch("agents.writer.chat_completion_with_retry", return_value=response(raw)):
            writer = get_writer_agent("test-key-not-a-credential")
            with self.assertRaisesRegex(ContentConstraintError, "页数"):
                writer(PPTGenerationState(topic="主题"), skeleton)

    def test_incomplete_or_empty_model_output_is_a_clear_failure(self):
        for text, reason, message in [('{"slides": [', "length", "被截断"), ('{broken}', "stop", "格式无效"), (None, "stop", "未返回完整内容")]:
            state = PPTGenerationState(topic="主题")
            result = SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content=text), finish_reason=reason)])
            with self.subTest(reason=reason, text=text), self.assertRaisesRegex(ContentConstraintError, message):
                parse_completion_payload(result, state, "Writer")
            self.assertEqual(len(state.messages), 1 if text else 0)

    def test_svg_replaces_only_an_empty_new_diagram_placeholder(self):
        slide = diagram_slide()
        skeleton = copy.deepcopy(slide); skeleton["diagram"] = None
        slide["diagram"] = {"nodes": [], "edges": []}
        slide["svg_code"] = '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 300 100"><text x="10" y="50">完整流程</text></svg>'
        normalized = normalize_outline_payload({"slides": [slide]}, skeleton={"slides": [skeleton]}, lock_scope="structure")
        outline = PPTOutline.model_validate(normalized)
        self.assertEqual(outline.slides[0].layout_type, "svg")
        self.assertIsNone(outline.slides[0].diagram)
        self.assertEqual(outline.slides[0].svg_code, slide["svg_code"])
        self.assertEqual(skeleton["layout_type"], "diagram")
        original = diagram_slide()
        preserved = normalize_outline_payload({"slides": [slide]}, skeleton={"slides": [original]}, lock_scope="structure")
        self.assertEqual(preserved["slides"][0]["diagram"], original["diagram"])
        slide["svg_code"] = None
        with self.assertRaises(ValidationError):
            PPTOutline.model_validate(normalize_outline_payload({"slides": [slide]}, skeleton={"slides": [skeleton]}, lock_scope="structure"))

    def test_content_text_unwraps_only_single_text_wrappers_and_preserves_locks(self):
        slide = copy.deepcopy(outline_fixture()["slides"][0])
        slide["bullet_points"] = ["一", "二", "三", "四"]
        slide["full_content"] = [{"text": "正文一"}, '{"text":"正文二"}', '{"text":"代码","id":1}', 'JSON 示例：{"text":"保留"}']
        result = normalize_outline_payload({"slides": [slide]})["slides"][0]
        self.assertEqual(result["full_content"], ["正文一", "正文二", '{"text":"代码","id":1}', 'JSON 示例：{"text":"保留"}'])
        old = copy.deepcopy(slide); old["full_content"] = ['{"text":"原文"}', "二", "三", "四"]
        protected = normalize_outline_payload({"slides": [old]}, skeleton={"slides": [old]}, lock_scope="design")["slides"][0]
        self.assertEqual(protected["full_content"], old["full_content"])

    def test_writer_repairs_a_complete_invalid_diagram_once_and_keeps_raw(self):
        invalid = diagram_slide(); invalid["diagram"]["nodes"][0]["shape_type"] = "cylinder"
        fixed = copy.deepcopy(invalid); fixed["diagram"] = None
        fixed["svg_code"] = '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 300 100"><text x="10" y="50">用户需求 → 内容规划 → 输出校验</text></svg>'
        skeleton = copy.deepcopy(invalid); skeleton["diagram"] = None
        state = PPTGenerationState(topic="主题")
        with patch("agents.writer.create_client"), patch("agents.writer.chat_completion_with_retry", side_effect=[response({"slides": [invalid]}), response({"slides": [fixed]})]) as calls:
            result = get_writer_agent("test-key-not-a-credential")(state, {"slides": [skeleton]})
        self.assertEqual(calls.call_count, 2)
        self.assertEqual(len(state.messages), 2)
        self.assertIn("cylinder", state.messages[0]["content"])
        self.assertEqual(result.slides[0].layout_type, "svg")
        with patch("agents.writer.create_client"), patch("agents.writer.chat_completion_with_retry", return_value=response({"slides": [invalid]})) as calls:
            with self.assertRaises(ValidationError):
                get_writer_agent("test-key-not-a-credential")(PPTGenerationState(topic="主题"), {"slides": [skeleton]})
        self.assertEqual(calls.call_count, 2)

    def test_layout_refs_restore_exact_content_and_structured_data(self):
        slide = table_slide()
        raw = {"canvases": [{"elements": [{"type": "textbox", "left": 1, "top": 0.5, "width": 11, "height": 0.9,
                                          "paragraphs": [{"content_ref": "title", "font_size": 32, "space_after": 0}]},
                                         {"type": "table", "left": 1, "top": 1.7, "width": 11, "height": 4.8}]}]}
        with patch("agents.template.create_client"), patch("agents.template.chat_completion_with_retry", return_value=response(raw)):
            blueprint = get_layout_agent("test-key-not-a-credential")(PPTGenerationState(topic="主题"), {"slides": [slide]})
        self.assertEqual(blueprint.canvases[0].elements[0].paragraphs[0].text, slide["title"])
        self.assertEqual(blueprint.canvases[0].elements[0].paragraphs[0].space_after, 0)
        self.assertEqual(blueprint.canvases[0].elements[1].table.model_dump(), slide["table"])
        self.assertFalse(blueprint.canvases[0].layout_resolved)

    def test_sanitizers_preserve_native_data_and_references_and_reject_fake_coverage(self):
        slide = table_slide()
        normalized = normalize_outline_payload({"schema_version": 2, "topic": "主题", "slides": [slide]})
        self.assertEqual(normalized["slides"][0]["table"], slide["table"])
        canvas = canvas_for(slide)
        normalized_canvas = normalize_blueprint_payload({"schema_version": 2, "canvases": [canvas.model_dump()]})
        self.assertEqual(normalized_canvas["canvases"][0]["elements"][0]["paragraphs"][0]["content_ref"], "title")
        payload = canvas.model_dump()
        payload["elements"][0]["paragraphs"][0]["space_after"] = None
        self.assertEqual(normalize_blueprint_payload({"canvases": [payload]})["canvases"][0]["elements"][0]["paragraphs"][0]["space_after"], 8)
        with self.assertRaises(ContentConstraintError):
            bind_canvas({"elements": [{"type": "shape", "left": 1, "top": 1, "width": 1, "height": 1,
                                       "paragraphs": [{"text": "指标对比", "content_ref": "title"}]}]}, slide)


if __name__ == "__main__":
    unittest.main()
