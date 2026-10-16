import copy
import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openai import APIConnectionError
from pptx import Presentation
import app as web
import main
from config.content import ContentConstraintError, identified_slide
from config.design import validate_icons, icon_suggestions, arrange_text_and_icons, validate_candidate
from config.jobs import execute_job
from config.preview import PreviewError
from config.schema import PPTFinalBlueprint, CanvasElement
from config.state import PPTGenerationState
from config.storage import PPTStore, content_hash
from test_content import table_slide, diagram_slide, canvas_for


class DesignTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="ppt_design_"); self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name); self.store = PPTStore(self.root)
        self.requests = []
        for replacement in (patch.object(web, "OUTPUT_ROOT", self.root), patch.object(web, "dispatch_pending"),
                            patch("config_manager.load_base_url", return_value=""),
                            patch.dict(os.environ, {"DASHSCOPE_API_KEY": "test-key", "PPT_DESIGN_MAX_LAYOUT_CALLS": "12"}),
                            patch("config.design.tool_signature", return_value={"test": "runtime"}),
                            patch("config.design.render_preview", side_effect=self.preview),
                            patch("config.design.get_layout_agent", return_value=self.designer)):
            replacement.start(); self.addCleanup(replacement.stop)
        self.client = web.app.test_client()
        slides = [identified_slide(table_slide(), 1), identified_slide(diagram_slide(), 2)]
        slides[1]["id"] = "slide-2"
        self.outline = {"schema_version": 2, "topic": "合成数据与流程", "slides": slides}
        self.blueprint = PPTFinalBlueprint(schema_version=2, canvases=[main.PPTXRenderer.resolve_canvas(canvas_for(slide)) for slide in slides])
        folder = self.root / "seed"; folder.mkdir(); ppt = folder / "original.pptx"
        main._render_blueprint(self.blueprint, ppt)
        state = PPTGenerationState(topic="合成数据与流程", requirements="保留数字和流程")
        state.set("outline", self.outline); state.set("blueprint", self.blueprint.model_dump()); state.set("last_file", str(ppt))
        request = {"user_input": "合成演示", "filename": str(ppt)}
        job, _ = self.store.create_job(request, state.snapshot(), folder, "seed-design-case", content_hash(request))
        self.store.claim_job(job["id"], "seed"); self.version = self.store.finish_job(job["id"], "seed", state.snapshot(), [self.store.describe_asset(ppt, "pptx")])
        self.seed_job_id = job["id"]; self.conversation = job["conversation_id"]
        self.client.post(f"/api/conversations/{self.conversation}/select")

    def preview(self, filename, directory, signature):
        directory = Path(directory); directory.mkdir(parents=True, exist_ok=True)
        pdf = directory / "result.pdf"; pdf.write_bytes(b"test PDF")
        pages = []
        for number in range(1, len(Presentation(filename).slides) + 1):
            path = directory / f"page-{number}.png"; path.write_bytes(b"test PNG"); pages.append(str(path))
        return {"pdf": str(pdf), "pages": pages, "pptx_sha256": content_hash(Path(filename).read_bytes().hex())}

    def designer(self, state, payload):
        self.requests.append(copy.deepcopy(payload))
        canvas = canvas_for(payload["slides"][0]); direction = payload["design_direction"]["id"]
        canvas.elements[0].top = .6 if direction == "clear" else .9
        component = next(element for element in canvas.elements if element.type in {"table", "diagram"})
        component.left = .9 if direction == "clear" else 1.2
        component.width = 11 if direction == "clear" else 10.5
        return PPTFinalBlueprint(schema_version=2, canvases=[canvas])

    def submit(self, key="candidate-request", pages=None):
        return self.client.post('/api/design/candidates', json={"conversation_id": self.conversation, "version_id": self.version,
            "request_id": key, "pages": pages if pages is not None else [1], "design_request": "使用对应语义图标，不生图"})

    def run_candidates(self):
        response = self.submit(); self.assertEqual(response.status_code, 202)
        execute_job(self.root, response.json["job_id"])
        return self.client.get(f'/api/conversations/{self.conversation}/design').json['jobs'][0]

    def adopt_payload(self, job, candidate=None):
        candidate = candidate or job['candidates'][0]
        return {"conversation_id": self.conversation, "version_id": self.version, "source_job_id": job['id'],
                "candidate_id": candidate['id'], "candidate_hash": next(item['sha256'] for item in candidate['assets'] if item['role']=='blueprint'), "request_id": "adopt-candidate-case"}

    def test_two_candidates_keep_semantics_other_page_and_current_pointer(self):
        with patch('main.run_ppt_system') as generate, patch('config.visual.run_visual_pipeline') as vision:
            job = self.run_candidates()
        generate.assert_not_called(); vision.assert_not_called()
        self.assertEqual(job['status'], 'succeeded'); self.assertEqual([item['status'] for item in job['candidates']], ['ready','ready'])
        self.assertEqual(self.store.conversation(self.conversation)['current_version_id'], self.version)
        self.assertEqual(len(self.store.versions(self.conversation)), 1)
        self.assertEqual(self.client.get('/state').json['job_id'], self.seed_job_id)
        for candidate in job['candidates']:
            asset = next(item for item in candidate['assets'] if item['role']=='blueprint')
            blueprint = self.client.get(asset['download_url']).json
            self.assertEqual(blueprint['canvases'][1], self.blueprint.canvases[1].model_dump())
            self.assertEqual(blueprint['canvases'][0]['elements'][1]['table']['rows'][0]['cells'][1]['value'], .8765)
        self.assertEqual(PPTStore(self.root).design_jobs(self.conversation)[0]['id'], job['id'])

    def test_adopt_is_idempotent_atomic_version_and_calls_no_model(self):
        job = self.run_candidates(); payload = self.adopt_payload(job)
        response = self.client.post('/api/design/adopt', json=payload); self.assertEqual(response.status_code,202)
        with patch('config.design.get_layout_agent') as model, patch('config.visual.run_visual_pipeline') as vision:
            execute_job(self.root, response.json['job_id'])
        model.assert_not_called(); vision.assert_not_called()
        duplicate = self.client.post('/api/design/adopt',json=payload)
        self.assertEqual(duplicate.json['job_id'], response.json['job_id'])
        versions = self.store.versions(self.conversation); self.assertEqual(len(versions),2)
        current = self.store.version(self.store.conversation(self.conversation)['current_version_id'])
        self.assertEqual(current['parent_id'], self.version)
        state = json.loads(current['snapshot_json'])['data']
        self.assertEqual(state['outline'], self.outline)
        self.assertEqual(state['blueprint']['canvases'][1], self.blueprint.canvases[1].model_dump())
        self.assertEqual(state['visual_review']['status'],'not_run')
        self.assertEqual(state['requested_mode'],'modify_design')

    def test_legacy_missing_outline_is_clear_error_before_enqueue(self):
        version=self.store.version(self.version); snapshot=json.loads(version['snapshot_json']);snapshot['data']['outline']=None
        with self.store.connection(write=True) as db:
            db.execute('UPDATE versions SET snapshot_json=? WHERE id=?',(json.dumps(snapshot,ensure_ascii=False),self.version))
        response=self.submit()
        self.assertEqual(response.status_code,400);self.assertIn('旧文稿',response.json['message'])
        self.assertFalse(self.store.design_jobs(self.conversation))

    def test_invalid_pages_and_owner_are_rejected_before_enqueue(self):
        for pages in ([True], [], [1,1], [1,2,3,4], [9]):self.assertEqual(self.submit('invalid-design',pages).status_code,400)
        with patch.object(web, '_owner',return_value='other'):
            self.assertEqual(self.submit().status_code,404)
            self.assertEqual(self.client.get(f'/api/conversations/{self.conversation}/design').status_code,404)

    def test_tampered_candidate_cannot_download_or_commit(self):
        job = self.run_candidates(); candidate=job['candidates'][0]
        asset=next(item for item in candidate['assets'] if item['role']=='pptx')
        stored = json.loads(self.store.job(job['id'])['snapshot_json'])['data']['design_candidates']['candidates'][0]['assets']
        path = self.store.asset_path(next(item['path'] for item in stored if item['id']==asset['id']));path.write_bytes(b'changed')
        self.assertEqual(self.client.get(asset['download_url']).status_code,409)
        self.assertEqual(self.client.post('/api/design/adopt',json=self.adopt_payload(job)).status_code,409)
        self.assertEqual(self.store.conversation(self.conversation)['current_version_id'],self.version)

    def test_new_current_version_prevents_old_candidate_adoption(self):
        job=self.run_candidates(); payload=self.adopt_payload(job)
        first=self.client.post('/api/design/adopt',json=payload);execute_job(self.root,first.json['job_id'])
        payload['candidate_id']=job['candidates'][1]['id'];payload['candidate_hash']=next(item['sha256'] for item in job['candidates'][1]['assets'] if item['role']=='blueprint');payload['request_id']='stale-candidate-case'
        self.assertEqual(self.client.post('/api/design/adopt',json=payload).status_code,409)

    def test_unknown_second_candidate_preserves_first_and_never_blind_retries(self):
        def designer(state,payload):
            if payload['design_direction']['id']=='visual':raise APIConnectionError(request=None)
            return self.designer(state,payload)
        with patch('config.design.get_layout_agent',return_value=designer):job=self.run_candidates()
        self.assertEqual(job['status'],'unknown');self.assertEqual(job['candidates'][0]['status'],'ready')
        self.assertEqual(self.client.post(f"/api/jobs/{job['id']}/retry",json={}).status_code,409)
        self.assertEqual(len(self.requests),1)
        self.assertEqual(self.store.conversation(self.conversation)['current_version_id'],self.version)

    def test_failed_screenshot_and_budget_are_explicit_without_commit(self):
        with patch('config.design.render_preview',side_effect=PreviewError('preview unavailable')):job=self.run_candidates()
        self.assertTrue(all(item['status']=='rejected' for item in job['candidates']))
        self.assertEqual(len(self.store.versions(self.conversation)),1)

    def test_changed_candidate_hash_rejected_before_commit(self):
        job=self.run_candidates();payload=self.adopt_payload(job);payload['candidate_hash']='not-current-hash'
        self.assertEqual(self.client.post('/api/design/adopt',json=payload).status_code,409)

    def test_unknown_request_still_consumes_budget_after_explicit_retry(self):
        with patch.dict(os.environ, {'PPT_DESIGN_MAX_LAYOUT_CALLS':'1'}),patch('config.design.get_layout_agent',return_value=lambda *args: (_ for _ in ()).throw(APIConnectionError(request=None))):
            job=self.run_candidates()
        self.assertEqual(job['status'],'unknown')
        self.store.retry_job(job['id'],confirm_unknown=True)
        with patch.dict(os.environ, {'PPT_DESIGN_MAX_LAYOUT_CALLS':'1'}):execute_job(self.root,job['id'])
        self.assertEqual(len(self.requests),0)
        saved=json.loads(self.store.job(job['id'])['snapshot_json'])['data']
        self.assertEqual(saved['design_call_attempts'],1)
        self.assertTrue(all('预算' in candidate['reason'] for candidate in saved['design_candidates']['candidates']))

    def test_known_budget_stops_extra_layout_requests(self):
        with patch.dict(os.environ, {'PPT_DESIGN_MAX_LAYOUT_CALLS':'1'}):job=self.run_candidates()
        self.assertEqual(len(self.requests),1)
        self.assertEqual([item['status'] for item in job['candidates']],['ready','rejected'])
        self.assertIn('预算',job['candidates'][1]['reason'])

    def test_explicit_theme_change_is_preserved_when_adopting_native_table(self):
        def dark_designer(state,payload):
            result=self.designer(state,payload);canvas=result.canvases[0];canvas.background_color='#101820'
            for e in canvas.elements:
                for p in e.paragraphs or []:p.font_color='#F8FAFC'
                if e.type=='table':
                    e.table_style.font_color='#F8FAFC';e.table_style.body_fill='#101820';e.table_style.alternate_fill='#223344'
            return result
        with patch('config.design.get_layout_agent',return_value=dark_designer):
            response=self.client.post('/api/design/candidates',json={'conversation_id':self.conversation,'version_id':self.version,
                'request_id':'dark-theme-candidate','pages':[1],'design_request':'深色背景、浅色文字，缩放文字以符合排版'})
            execute_job(self.root,response.json['job_id'])
        job=self.client.get(f'/api/conversations/{self.conversation}/design').json['jobs'][0]
        self.assertEqual(job['candidates'][0]['status'],'ready')
        result=self.client.post('/api/design/adopt',json=self.adopt_payload(job));execute_job(self.root,result.json['job_id'])
        latest=self.store.version(self.store.conversation(self.conversation)['current_version_id'])
        data=json.loads(latest['snapshot_json'])['data']
        self.assertEqual(data['blueprint']['canvases'][0]['background_color'],'#101820')
        self.assertEqual(data['outline']['slides'][0]['theme']['background'],'#101820')
        native=next(e for e in data['blueprint']['canvases'][0]['elements'] if e['type']=='table')
        self.assertEqual(native['table_style']['font_color'],'#F8FAFC')
        self.assertEqual(data['blueprint']['canvases'][1],self.blueprint.canvases[1].model_dump())


class DesignIconTests(unittest.TestCase):
    def test_real_icon_requires_valid_library_name_content_anchor_and_alignment(self):
        slide = table_slide(); canvas = canvas_for(slide)
        icon = CanvasElement(type='icon',id='icon-1',content_ref='title',icon_name='coffee',left=.25,top=.7,width=.25,height=.25)
        canvas.elements.append(icon); validate_icons(canvas,slide)
        icon.icon_name='circle'
        with self.assertRaises(ContentConstraintError):validate_icons(canvas,slide)
        icon.icon_name='coffee';icon.content_ref='unknown'
        with self.assertRaises(ContentConstraintError):validate_icons(canvas,slide)
        icon.content_ref='title';icon.top=3
        with self.assertRaises(ContentConstraintError):validate_icons(canvas,slide)

    def test_candidate_reflow_uses_saved_coordinates_and_moves_icons_with_text(self):
        from renderer import PPTXRenderer
        from config.schema import SlideCanvas
        slide = table_slide(); slide.update(layout_type="standard_text", table=None,
            bullet_points=["全球咖啡种植带的地理分布与风土特点" for _ in range(4)],
            full_content=["必须保留这段较长的原始说明以及87.65%的数字和单位，不能通过删改文字或缩小字号让它塞进容器。" for _ in range(4)],
            bullet_ids=[f"point-{i}" for i in range(1,5)], punchline="保留原始全部内容")
        elements=[{"type":"textbox","id":"title","left":.8,"top":.4,"width":11.7,"height":.9,"paragraphs":[{"text":slide['title'],"content_ref":"title","font_size":32}]},
                  {"type":"textbox","id":"punchline","left":.8,"top":1.3,"width":11.7,"height":.6,"paragraphs":[{"text":slide['punchline'],"content_ref":"punchline","font_size":18}]}]
        for i in range(4):
            x=1 if i%2==0 else 7; top=2.4 if i<2 else 4.9
            elements.extend([{"type":"icon","id":f"icon-{i}","icon_name":"sprout","left":x,"top":top,"width":.4,"height":.4},
               {"type":"textbox","id":f"head-{i}","left":x+.6,"top":top,"width":4.8,"height":.4,"paragraphs":[{"text":slide['bullet_points'][i],"content_ref":f"bullet_points/point-{i+1}","font_size":20}]},
               {"type":"textbox","id":f"body-{i}","left":x,"top":top+.5,"width":5.4,"height":1.4,"paragraphs":[{"text":slide['full_content'][i],"content_ref":f"full_content/point-{i+1}","font_size":18}]}])
        original=SlideCanvas.model_validate({"slide_id":slide['id'],"elements":elements})
        saved=original.model_dump();resolved=PPTXRenderer.resolve_canvas(arrange_text_and_icons(original),fit_text=True)
        validate_icons(resolved,slide);self.assertEqual(original.model_dump(),saved)
        self.assertEqual([(p.text,p.font_size) for e in original.elements for p in e.paragraphs or []],[(p.text,p.font_size) for e in resolved.elements for p in e.paragraphs or []])
        self.assertTrue(all(e.content_ref for e in resolved.elements if e.type=='icon'))

    def test_alternative_reading_axis_keeps_all_words_and_readable_font_size(self):
        from renderer import PPTXRenderer
        slide=table_slide();slide.update(layout_type="standard_text",table=None,full_content=["原文数据87.65%必须保留。"],bullet_points=["原要点"],bullet_ids=["point-1"])
        canvas=canvas_for(slide)
        paragraph=next(p for e in canvas.elements for p in e.paragraphs or [] if (p.content_ref or "").startswith("full_content/"));paragraph.font_size=14
        result=PPTXRenderer.resolve_canvas(arrange_text_and_icons(canvas,{"full_content/point-1":20},flow_text=True,slide=slide),fit_text=True)
        copied=next(p for e in result.elements for p in e.paragraphs or [] if p.content_ref==paragraph.content_ref)
        self.assertEqual(copied.text,paragraph.text);self.assertEqual(copied.font_size,20)
        self.assertGreater(result.elements[1].width,canvas.elements[1].width)

    def test_optional_visible_bullet_cannot_disappear_from_candidate(self):
        from renderer import PPTXRenderer
        slide=table_slide();slide.update(layout_type="standard_text",table=None,full_content=["必须保留87.65%。"],bullet_points=["不可遗漏的要点标题"],bullet_ids=["point-1"])
        candidate=PPTXRenderer.resolve_canvas(canvas_for(slide))
        baseline=candidate.model_copy(deep=True)
        baseline.elements.append(CanvasElement(type="textbox",id="visible-heading",left=.7,top=3.2,width=10,height=.7,
            paragraphs=[{"content_ref":"bullet_points/point-1","text":slide['bullet_points'][0],"font_size":20}]))
        with self.assertRaisesRegex(ContentConstraintError,"遗漏"):
            validate_candidate(PPTFinalBlueprint(canvases=[candidate]),PPTFinalBlueprint(canvases=[baseline]),{"slides":[slide]},[1])

    def test_coffee_suggestions_are_whitelisted_and_bound_to_original_text(self):
        slide=table_slide();slide['full_content']=['幼苗种植与采摘','烘焙加热','冲煮萃取后饮用']
        suggestions=icon_suggestions(slide)
        self.assertEqual([item['icon_name'] for item in suggestions],['sprout','flame','droplet'])
        self.assertTrue(all(item['content_ref'].startswith('full_content/') for item in suggestions))


if __name__=='__main__':unittest.main()
