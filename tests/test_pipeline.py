import copy
import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from pptx import Presentation

import main
from config.schema import PPTFinalBlueprint, PPTOutline
from config.state import PPTGenerationState
from config.llm_client import resolve_model_name


def outline_fixture():
    return {"topic": "固定内容", "slides": [{
        "title": f"第 {number} 页", "layout_type": "standard_text", "bullet_points": ["原要点"],
        "full_content": ["原正文"], "theme": {
            "primary": "#005088", "secondary": "#11CAA0", "content": "#334155", "background": "#FFFFFF",
        },
    } for number in range(1, 4)]}


def blueprint_fixture():
    return {"canvases": [{"background_color": "#FFFFFF", "elements": [{
        "type": "textbox", "left": 1, "top": 1, "width": 6, "height": 1,
        "paragraphs": [{"content_ref": "title", "text": f"第 {number} 页"}],
    }, {
        "type": "textbox", "left": 1, "top": 2.2, "width": 6, "height": 1,
        "paragraphs": [{"content_ref": "full_content/point-1", "text": "原正文"}],
    }]} for number in range(1, 4)]}


class ImageEndpointConfigurationTests(unittest.TestCase):
    def test_image_endpoint_defaults_to_public_api_without_a_network_call(self):
        from agents import image_generator
        with patch.dict(os.environ, {"DASHSCOPE_IMAGE_BASE_URL": ""}), patch.object(
            image_generator.dashscope, "api_key", ""
        ), patch.object(image_generator.dashscope, "base_http_api_url", "https://previous.example.invalid"), patch.object(
            image_generator.MultiModalConversation, "call"
        ) as request:
            generator = image_generator.get_image_generator("test-key-not-a-credential")
            self.assertTrue(callable(generator))
            self.assertEqual(image_generator.dashscope.base_http_api_url, "https://dashscope.aliyuncs.com/api/v1")
            request.assert_not_called()

    def test_image_endpoint_uses_native_override_independently_of_text_gateway(self):
        from agents import image_generator
        with patch.dict(os.environ, {"DASHSCOPE_IMAGE_BASE_URL": "  https://image.example.invalid/api/v1  ",
                                     "DASHSCOPE_BASE_URL": "https://text.example.invalid/compatible-mode/v1"}), patch.object(
            image_generator.dashscope, "api_key", ""
        ), patch.object(image_generator.dashscope, "base_http_api_url", ""), patch.object(
            image_generator.MultiModalConversation, "call"
        ) as request:
            image_generator.get_image_generator("test-key-not-a-credential")
            self.assertEqual(image_generator.dashscope.base_http_api_url, "https://image.example.invalid/api/v1")
            request.assert_not_called()


class PipelineRegressionTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix="ppt_pipeline_test_")
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.output = self.root / "result.pptx"
        self.state = PPTGenerationState(topic="固定内容")
        self.state.set("outline", outline_fixture())
        self.state.set("blueprint", blueprint_fixture())
        replacement = patch.dict(os.environ, {"DASHSCOPE_API_KEY": "test-key-not-a-credential"})
        replacement.start()
        self.addCleanup(replacement.stop)

    def call_incremental(self, pages):
        return main.run_incremental_modify_pages(
            "修改第 2 页", pages, filename=str(self.output), state=self.state, run_dir=str(self.root)
        )

    def test_invalid_targets_never_fall_back_to_full_generation(self):
        for pages in ([0], [4], ["2"], [True]):
            with self.subTest(pages=pages), patch.object(main, "run_ppt_system") as full, patch.object(
                main, "get_writer_agent"
            ) as writer:
                self.assertIsNone(self.call_incremental(pages))
                full.assert_not_called()
                writer.assert_not_called()
                self.assertFalse(self.output.exists())

    def test_missing_cache_never_falls_back_to_full_generation(self):
        self.state.set("blueprint", None)
        with patch.object(main, "run_ppt_system") as full:
            self.assertIsNone(self.call_incremental([2]))
            full.assert_not_called()

    def test_missing_key_never_turns_page_edit_into_demo(self):
        with patch.dict(os.environ, {"DASHSCOPE_API_KEY": ""}), patch.object(main, "run_ppt_system") as full:
            self.assertIsNone(self.call_incremental([2]))
            full.assert_not_called()
        self.assertFalse(self.output.exists())

    def test_failed_page_edit_restores_both_caches(self):
        original_outline = copy.deepcopy(self.state.get("outline"))
        original_blueprint = copy.deepcopy(self.state.get("blueprint"))
        def fail_writer(state, payload):
            state.set("outline", {"topic": "意外覆盖", "slides": []})
            raise RuntimeError("模拟 Writer 故障")
        with patch.object(main, "get_writer_agent", return_value=fail_writer), patch.object(
            main, "get_image_generator"
        ), patch.object(main, "get_layout_agent"), patch.object(main, "run_ppt_system") as full:
            self.assertIsNone(self.call_incremental([2]))
            full.assert_not_called()
        self.assertEqual(self.state.get("outline"), original_outline)
        self.assertEqual(self.state.get("blueprint"), original_blueprint)
        self.assertEqual(self.state.get("generation_status"), "failed")

    def test_successful_page_edit_preserves_other_pages(self):
        original_outline = copy.deepcopy(self.state.get("outline"))
        original_blueprint = copy.deepcopy(self.state.get("blueprint"))
        def writer(state, payload):
            changed = copy.deepcopy(payload)
            changed["slides"][0]["title"] = "修改后的第 2 页"
            return PPTOutline.model_validate(changed)
        def layout(state, payload):
            changed = copy.deepcopy(blueprint_fixture()["canvases"][1])
            changed["elements"][0]["paragraphs"][0]["text"] = payload["slides"][0]["title"]
            return PPTFinalBlueprint.model_validate({"canvases": [changed]})
        with patch.object(main, "get_writer_agent", return_value=writer) as writer_factory, patch.object(
            main, "get_layout_agent", return_value=layout
        ) as layout_factory, patch.object(main, "get_image_generator") as image:
            self.assertEqual(self.call_incremental([2]), str(self.output))
        self.assertEqual(writer_factory.call_args.kwargs["model_name"], resolve_model_name("writer"))
        self.assertEqual(layout_factory.call_args.kwargs["model_name"], resolve_model_name("layout"))
        self.assertEqual(image.call_args.kwargs["model_name"], resolve_model_name("image"))
        self.assertEqual(len(Presentation(self.output).slides), 3)
        image.return_value.assert_not_called()
        self.assertEqual(self.state.get("outline")["slides"][1]["title"], "修改后的第 2 页")
        for index in (0, 2):
            self.assertEqual(self.state.get("outline")["slides"][index], original_outline["slides"][index])
            self.assertEqual(self.state.get("blueprint")["canvases"][index], original_blueprint["canvases"][index])

    def test_cloud_content_failure_does_not_generate_demo(self):
        with patch.object(main, "get_planner_agent", side_effect=RuntimeError("模拟模型故障")), patch.object(
            main, "_render_offline_demo"
        ) as demo:
            result = main.run_ppt_system("主题", filename=str(self.output), state=self.state, run_dir=str(self.root))
        self.assertIsNone(result)
        demo.assert_not_called()
        self.assertEqual(self.state.get("generation_status"), "failed")

    def test_layout_failure_does_not_generate_demo(self):
        source = self.root / "content.json"
        source.write_text(json.dumps(outline_fixture()), encoding="utf-8")
        with patch.object(main, "get_layout_agent", side_effect=RuntimeError("模拟 Layout 故障")), patch.object(
            main, "_render_offline_demo"
        ) as demo:
            result = main.run_ppt_system(
                "主题", filename=str(self.output), from_content_json=str(source), state=self.state, run_dir=str(self.root)
            )
        self.assertIsNone(result)
        demo.assert_not_called()
        self.assertEqual(self.state.get("generation_status"), "failed")

    def test_explicit_offline_generation_is_marked_as_demo(self):
        with patch.dict(os.environ, {"DASHSCOPE_API_KEY": ""}):
            result = main.run_ppt_system("离线主题", filename=str(self.output), state=self.state, run_dir=str(self.root))
        self.assertEqual(result, str(self.output))
        self.assertEqual(len(Presentation(self.output).slides), 3)
        self.assertEqual(self.state.get("result_kind"), "offline_demo")
        presentation = Presentation(self.output)
        for slide in presentation.slides:
            for shape in slide.shapes:
                self.assertLessEqual(shape.left + shape.width, presentation.slide_width)
                self.assertLessEqual(shape.top + shape.height, presentation.slide_height)

    def test_model_defaults_and_explicit_environment_override(self):
        roles = {"planner": "qwen3.8-27b", "writer": "qwen3.8-27b", "layout": "qwen3.8-27b", "image": "qwen-image-2.1-pro"}
        with patch.dict(os.environ, {f"PPT_{role.upper()}_MODEL": "" for role in roles}):
            for role, expected in roles.items():
                self.assertEqual(resolve_model_name(role), expected)
            with patch.dict(os.environ, {"PPT_WRITER_MODEL": " glm-5.3 "}):
                self.assertEqual(resolve_model_name("writer"), "glm-5.3")

    def test_failed_content_review_is_explicitly_marked(self):
        content = PPTOutline.model_validate(outline_fixture())
        page = PPTFinalBlueprint.model_validate({"canvases": [blueprint_fixture()["canvases"][0]]})
        with patch.object(main, "get_planner_agent", return_value=lambda *args, **kwargs: content), patch.object(
            main, "get_writer_agent", return_value=lambda *args, **kwargs: content
        ), patch.object(main, "get_image_generator"), patch.object(
            main, "get_layout_agent", return_value=lambda *args, **kwargs: page
        ), patch.object(main, "review_ppt_outline", return_value=(False, ["模拟审核不通过"])):
            result = main.run_ppt_system("主题", filename=str(self.output), state=self.state, run_dir=str(self.root))
        self.assertEqual(result, str(self.output))
        self.assertEqual(self.state.get("content_review_status"), "needs_review")

    def test_missing_requested_image_is_a_failure(self):
        content = PPTOutline.model_validate(outline_fixture())
        content.slides[0].image_prompt = "需要配图"
        original_outline = copy.deepcopy(self.state.get("outline"))
        with patch.object(main, "get_writer_agent", return_value=lambda *args: PPTOutline(
            topic=content.topic, slides=[content.slides[0]]
        )), patch.object(main, "get_image_generator", return_value=lambda **kwargs: None), patch.object(
            main, "get_layout_agent"
        ) as layout:
            self.assertIsNone(self.call_incremental([2]))
        layout.return_value.assert_not_called()
        self.assertEqual(self.state.get("outline"), original_outline)
        self.assertFalse(self.output.exists())

    def test_blueprint_replay_never_calls_agents(self):
        source = self.root / "blueprint.json"
        source.write_text(json.dumps(blueprint_fixture()), encoding="utf-8")
        with patch.object(main, "get_planner_agent") as planner, patch.object(main, "get_layout_agent") as layout:
            result = main.run_ppt_system(
                "缓存回放", filename=str(self.output), from_blueprint_json=str(source), state=self.state, run_dir=str(self.root)
            )
        self.assertEqual(result, str(self.output))
        planner.assert_not_called()
        layout.assert_not_called()
        self.assertEqual(len(Presentation(self.output).slides), 3)


if __name__ == "__main__":
    unittest.main()
