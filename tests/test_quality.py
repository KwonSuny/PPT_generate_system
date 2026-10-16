import copy
import json
import os
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from pptx import Presentation
import main
from agents.writer import get_writer_agent
from config.content import identified_slide, bind_canvas
from config.design import validate_candidate
from config.layout_budget import layout_issues, design_request_policy, font_floor
from config.llm_client import model_request_options
from config.schema import CanvasElement, PPTFinalBlueprint, SlideCanvas
from config.state import PPTGenerationState
from renderer import PPTXRenderer
from test_content import outline_fixture, canvas_for, table_slide


class QualityTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory(prefix='ppt_quality_');self.addCleanup(self.temp.cleanup)
        self.root=Path(self.temp.name)
        self.slide=identified_slide(copy.deepcopy(outline_fixture()['slides'][0]))

    def test_zero_height_separator_is_a_valid_line_and_preserved_in_export(self):
        canvas=canvas_for(self.slide)
        canvas.elements.append(CanvasElement(id='separator',type='shape',shape_type='line',left=.6,top=1.5,width=12.133,height=0))
        result=PPTXRenderer.resolve_canvas(canvas)
        self.assertTrue(any(e.id=='separator' for e in result.elements))
        self.assertFalse(any(i['kind']=='bounds' for i in layout_issues(result)))
        file=self.root/'line.pptx';main._render_blueprint(PPTFinalBlueprint(canvases=[result]),str(file))
        self.assertEqual(sum(shape.shape_type==9 for shape in Presentation(file).slides[0].shapes),1)

    def test_color_is_checked_against_the_actual_panel_and_survives_export(self):
        canvas=canvas_for(self.slide);title=canvas.elements[0]
        plate=CanvasElement(id='dark-panel',type='shape',shape_type='rectangle',left=title.left,top=title.top,
            width=title.width,height=title.height,fill_color='#101820')
        canvas.elements.insert(0,plate)
        title.paragraphs[0].font_color='#333333';title.paragraphs[0].font_family='SimHei'
        result=PPTXRenderer.resolve_canvas(canvas)
        para=result.elements[1].paragraphs[0]
        self.assertGreaterEqual(PPTXRenderer.__new__(PPTXRenderer)._contrast_ratio(para.font_color,'#101820'),4.5)
        file=self.root/'panel.pptx';main._render_blueprint(PPTFinalBlueprint(canvases=[result]),str(file))
        text=next(shape for shape in Presentation(file).slides[0].shapes if shape.has_text_frame and shape.text==self.slide['title'])
        run=text.text_frame.paragraphs[0]
        self.assertEqual(str(run.font.color.rgb),para.font_color.lstrip('#'))
        self.assertEqual(run.font.name,'SimHei')

    def test_readable_accent_keeps_its_hue_instead_of_forcing_black(self):
        r=PPTXRenderer.__new__(PPTXRenderer);color=r._safe_text_color('#FF6833','#FFFFFF')
        self.assertNotEqual(color,'#111827');self.assertGreaterEqual(r._contrast_ratio(color,'#FFFFFF'),4.5)

    def test_duplicate_display_is_optional_but_unique_content_is_protected(self):
        canvas=PPTXRenderer.resolve_canvas(canvas_for(self.slide));base=canvas.model_copy(deep=True)
        duplicate=base.elements[0].model_copy(deep=True);duplicate.id='duplicate-title';duplicate.top=4
        base.elements.append(duplicate)
        validate_candidate(PPTFinalBlueprint(canvases=[canvas]),PPTFinalBlueprint(canvases=[base]),{'slides':[self.slide]},[1])
        base.elements[-1].paragraphs[0].text='不同的独有信息'
        with self.assertRaisesRegex(Exception,'遗漏'):
            validate_candidate(PPTFinalBlueprint(canvases=[canvas]),PPTFinalBlueprint(canvases=[base]),{'slides':[self.slide]},[1])

    def test_font_resize_is_bounded_and_preserves_text(self):
        canvas=canvas_for(self.slide);title=canvas.elements[0];title.height=.7;title.paragraphs[0].font_size=40
        fitted=PPTXRenderer.fit_readable_text(canvas,{'title':28})
        self.assertEqual(fitted.elements[0].paragraphs[0].text,self.slide['title'])
        self.assertGreaterEqual(fitted.elements[0].paragraphs[0].font_size,28)
        self.assertLess(fitted.elements[0].paragraphs[0].font_size,40)
        policy=design_request_policy('缩放文字以符合排版，同时避免背景对文字显示造成影响。')
        self.assertTrue(policy['allow_font_resize']);self.assertTrue(policy['allow_theme_change'])
        self.assertFalse(design_request_policy('保持字号不变')['allow_font_resize'])
        self.assertFalse(design_request_policy('不改变字体大小，仅修改颜色')['allow_font_resize'])

    def test_thinking_options_do_not_disable_reasoning_only_models(self):
        with patch.dict(os.environ,{'PPT_WRITER_THINKING':'auto','PPT_LAYOUT_THINKING':'on'}):
            self.assertNotIn('enable_thinking',model_request_options('writer','kimi-k3')['extra_body'])
            self.assertNotIn('thinking_budget',model_request_options('writer','kimi-k3')['extra_body'])
            self.assertNotIn('temperature',model_request_options('writer','kimi-k3'))
            self.assertTrue(model_request_options('layout','qwen3.8-max')['extra_body']['enable_thinking'])
            self.assertEqual(model_request_options('layout','unlisted-model')['extra_body']['enable_thinking'],True)

    def test_writer_matches_requested_page_id_without_rewriting_other_pages(self):
        slide={**self.slide,'id':'slide-2'};extra={**copy.deepcopy(slide),'id':'slide-1','title':'范围外'}
        payload={'topic':'测试','slides':[extra,slide]}
        reply=SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content=json.dumps(payload,ensure_ascii=False)))])
        state=PPTGenerationState(topic='测试');state.set_operation('continue',scope='design',pages=[2])
        with patch('agents.writer.create_client'),patch('agents.writer.chat_completion_with_retry',return_value=reply):
            result=get_writer_agent('test-key')(state,{'topic':'测试','slides':[slide]})
        self.assertEqual([s.id for s in result.slides],['slide-2'])
        self.assertEqual(result.slides[0].title,slide['title'])

    def test_font_range_uses_lower_bound_and_partial_panel_is_reported(self):
        self.assertEqual(font_floor('title','标题28–36pt，正文至少18pt'),28)
        self.assertEqual(font_floor('full_content/point-1','正文18-24pt'),18)
        canvas=canvas_for(self.slide);title=canvas.elements[0]
        canvas.elements.insert(0,CanvasElement(id='half',type='shape',shape_type='rectangle',left=title.left,top=title.top,
            width=title.width/2,height=title.height,fill_color='#101820'))
        result=PPTXRenderer.resolve_canvas(canvas)
        self.assertTrue(any(i['kind']=='contrast' for i in layout_issues(result)))

    def test_completed_schema_error_enters_bounded_layout_repair(self):
        calls=[]
        def designer(state,payload):
            calls.append(payload)
            if len(calls)==1:
                CanvasElement(type='unknown',left=0,top=0,width=1,height=1)
            return PPTFinalBlueprint(canvases=[canvas_for(self.slide)])
        result=main._layout_page(PPTGenerationState(topic='验证'),designer,{'slides':[self.slide]},1,str(self.root))
        self.assertEqual(len(calls),2)
        self.assertIn('layout_feedback',calls[1]);self.assertTrue(result.layout_resolved)

    def test_four_notes_fallback_reserves_readable_svg_space(self):
        from config.layout_fallback import stable_layout
        from config.layout_budget import svg_display
        slide=copy.deepcopy(self.slide)
        slide.update(title='一杯咖啡的旅程',punchline='从种子到杯中的四段关键历程',layout_type='svg',
            bullet_ids=[f'point-{i}' for i in range(1,5)],
            bullet_points=['种植：培育优质咖啡树','采收处理：获取并加工生豆','烘焙：激发风味化学反应','研磨冲煮：萃取最终饮品'],
            full_content=['采收成熟果实后去除果肉并干燥，得到可储存运输的咖啡生豆。']*4,
            svg_code='<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 960 300"><rect width="960" height="300" fill="#FFF7E8"/><text x="50" y="150" font-size="17">完整的工艺流程</text></svg>')
        canvas,name=stable_layout(slide)
        self.assertEqual(name,'art-top-notes-4')
        self.assertGreaterEqual(svg_display(next(e for e in canvas.elements if e.type=='svg'))['estimated_min_label_pt'],10)
        self.assertEqual(sum(p.content_ref.startswith('full_content/') for e in canvas.elements for p in e.paragraphs or []),4)

    def test_native_style_defaults_follow_theme_and_source_uses_slide_background(self):
        slide=table_slide();slide['theme'].update(background='#12352A',content='#FFF7E8',primary='#D4AF37')
        canvas=bind_canvas({'background_color':'#12352A','elements':[
            {'type':'textbox','left':.6,'top':.4,'width':12,'height':1,'paragraphs':[{'content_ref':'title','text':'指标对比','font_size':32}]},
            {'id':'table','type':'table','left':.6,'top':2,'width':12,'height':3.8}]},slide,preserve_component_style=True)
        self.assertEqual(canvas.elements[-1].table_style.header_fill,'#D4AF37')
        canvas.elements[-1].table_style.body_fill='#FFFFFF';canvas.elements[-1].table_style.font_color='#334155'
        canvas=PPTXRenderer.resolve_canvas(canvas)
        path=self.root/'source.pptx';main._render_blueprint(PPTFinalBlueprint(canvases=[canvas]),str(path))
        source=next(shape for shape in Presentation(path).slides[0].shapes if shape.name=='table-source')
        color='#'+str(source.text_frame.paragraphs[0].font.color.rgb)
        self.assertGreaterEqual(PPTXRenderer.__new__(PPTXRenderer)._contrast_ratio(color,'#12352A'),4.5)

    def test_layout_deduplicates_native_source_without_inventing_a_body_font_floor(self):
        slide=table_slide();canvas=canvas_for(slide)
        canvas.elements.append(CanvasElement(id='duplicate-source',type='textbox',left=.6,top=6.7,width=12,height=.3,
            paragraphs=[{'content_ref':'table/source','text':slide['table']['source'],'font_size':12}]))
        result=main._layout_page(PPTGenerationState(topic='验证'),lambda *args:PPTFinalBlueprint(canvases=[canvas]),
            {'slides':[slide],'required_content_refs':['table/source']},1,str(self.root))
        self.assertFalse(any(e.id=='duplicate-source' for e in result.elements))
        self.assertEqual(next(e.table.source for e in result.elements if e.type=='table'),slide['table']['source'])
        self.assertEqual(font_floor('table/source','正文至少18pt，来源12pt'),12)

    def test_first_layout_also_enforces_the_requested_title_floor(self):
        calls=[]
        def designer(state,payload):
            calls.append(payload)
            canvas=canvas_for(self.slide)
            canvas.elements[0].paragraphs[0].font_size=20 if len(calls)==1 else 32
            return PPTFinalBlueprint(canvases=[canvas])
        result=main._layout_page(PPTGenerationState(topic='验证',requirements='标题28–36pt，正文至少18pt'),
            designer,{'slides':[self.slide]},1,str(self.root))
        self.assertEqual(len(calls),2)
        self.assertGreaterEqual(result.elements[0].paragraphs[0].font_size,28)
