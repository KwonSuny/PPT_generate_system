# -*- coding: utf-8 -*-
"""
Flask 层路由测试：验证增量条件满足时走 incremental_page 分支，并落盘资产快照。
不依赖真实 LLM API（无 DASHSCOPE_API_KEY 时增量管线回退离线演示）。
"""
import json
import sys
import time
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
# 确保无 API key（走离线回退，不产生真实网络调用）
import os
os.environ.pop("DASHSCOPE_API_KEY", None)

import app as app_module
from config.state import PPTGenerationState

from pptx import Presentation

# 隔离：重定向输出根目录与配置目录到临时目录，避免污染用户真实数据
import tempfile
_tmp = Path(tempfile.mkdtemp(prefix="ppt_routing_test_"))
app_module.OUTPUT_ROOT = _tmp / "output"
app_module.OUTPUT_ROOT.mkdir(parents=True, exist_ok=True)
import config_manager
config_manager.CONFIG_DIR = _tmp / "config"
config_manager.CONFIG_FILE = config_manager.CONFIG_DIR / "config.json"


def poll_state(timeout=60, expect_snapshot=False):
    def _disk():
        rd = app_module._active_run_dir()
        return rd / "state.json" if rd else None

    deadline = time.time() + timeout
    while time.time() < deadline:
        disk = _disk()
        if disk and disk.exists():
            try:
                with open(disk, "r", encoding="utf-8") as f:
                    payload = json.load(f)
                if not expect_snapshot or (payload.get("history_snapshots") or []):
                    return payload
            except Exception:
                pass
        time.sleep(0.5)
    return None


def test_parse_page_numbers():
    print("=" * 60)
    print("TEST 5: 页码解析（单页/多页/范围/混合）")
    print("=" * 60)
    f = app_module._parse_page_numbers
    assert f("2") == [2]
    assert f("2,3") == [2, 3]
    assert f("2-4") == [2, 3, 4]
    assert f("1,3-5") == [1, 3, 4, 5]
    assert f("3,1") == [1, 3], "应去重排序"
    assert f("") is None
    assert f("  ") is None
    assert f("abc") is None
    assert f("2,x") is None
    print("PASS: 页码解析正确\n")


def main():
    print("=" * 60)
    print("TEST 4: Flask 增量路由 + 资产落盘")
    print("=" * 60)
    client = app_module.app.test_client()

    # 1) 首页可达
    resp = client.get("/")
    assert resp.status_code == 200, f"GET / 应返回 200，实际 {resp.status_code}"
    print("PASS: GET / 返回 200")

    # 2) 预置缓存状态（模拟已有一次成功生成）
    outline = {
        "topic": "测试主题",
        "slides": [
            {"title": "第1页标题", "punchline": "一", "bullet_points": ["a"], "full_content": ["a"],
             "layout_type": "standard_text", "theme": {"primary": "#1E3A8A", "secondary": "#10B981",
             "content": "#64748B", "background": "#F8FAFC"}, "svg_code": None, "image_prompt": None,
             "image_path": None, "icons": []},
            {"title": "第2页标题", "punchline": "二", "bullet_points": ["b"], "full_content": ["b"],
             "layout_type": "standard_text", "theme": {"primary": "#1E3A8A", "secondary": "#10B981",
             "content": "#64748B", "background": "#F8FAFC"}, "svg_code": None, "image_prompt": None,
             "image_path": None, "icons": []},
        ],
    }
    seed = PPTGenerationState(topic="测试主题", requirements="初始")
    seed.set("outline", outline)
    app_module._save_state(seed)

    # 3) 携带 modify_page 的修改请求 → 应路由到单页增量
    resp = client.post("/generate", data={
        "topic": "测试主题",
        "requirements": "只改第2页的标题",
        "filename": "incremental_route.pptx",
        "mode": "modify_text",
        "modify_page": "2",
    })
    assert resp.status_code == 200, f"POST /generate 应返回 200，实际 {resp.status_code}"
    body = resp.get_data(as_text=True)
    assert "单页仅改文本" in body, "页面应显示单页仅改文本标签"
    print("PASS: POST /generate 路由到单页增量模式（页面标签正确）")

    # 4) 等待后台线程完成，校验状态与资产快照
    final = poll_state(expect_snapshot=True)
    assert final is not None, "state.json 应被后台线程写入"
    assert final.get("operation_mode") == "incremental_page", (
        f"operation_mode 应为 incremental_page，实际 {final.get('operation_mode')}"
    )
    print("PASS: state.operation_mode == incremental_page（增量路由生效）")

    snapshots = final.get("history_snapshots") or []
    assert snapshots, "应产生历史快照"
    last = snapshots[-1]
    assert last.get("mode") == "incremental_page"
    assert last.get("modify_page") == 2
    assert last.get("outline_path") and Path(last["outline_path"]).exists(), (
        f"快照应包含有效 outline_path，实际 {last.get('outline_path')}"
    )
    if last.get("blueprint_path"):
        assert Path(last["blueprint_path"]).exists(), "blueprint_path 若存在则必须有效"
    else:
        print("NOTE: 离线回退模式不产生 blueprint，blueprint_path 为空（真实 API 管线下会有）")
    print(f"PASS: 历史快照含资产路径 outline={last['outline_path']}")

    last_file = final.get("data", {}).get("last_file") or final.get("last_file")
    assert last_file and Path(last_file).exists(), f"应产出 pptx 文件，实际 {last_file}"
    prs = Presentation(last_file)
    assert len(prs.slides) >= 1
    print(f"PASS: 输出文件 {Path(last_file).name} 存在（{len(prs.slides)} 页，离线回退模式）")

    print("\nALL FLASK ROUTING TESTS PASSED")


def test_global_routing():
    print("=" * 60)
    print("TEST 6: modify_text 无页码 → 全局增量路由")
    print("=" * 60)
    client = app_module.app.test_client()
    outline = {
        "topic": "测试主题",
        "slides": [
            {"title": "第1页标题", "punchline": "一", "bullet_points": ["a"], "full_content": ["a"],
             "layout_type": "standard_text", "theme": {"primary": "#1E3A8A", "secondary": "#10B981",
             "content": "#64748B", "background": "#F8FAFC"}, "svg_code": None, "image_prompt": None,
             "image_path": None, "icons": []},
            {"title": "第2页标题", "punchline": "二", "bullet_points": ["b"], "full_content": ["b"],
             "layout_type": "standard_text", "theme": {"primary": "#1E3A8A", "secondary": "#10B981",
             "content": "#64748B", "background": "#F8FAFC"}, "svg_code": None, "image_prompt": None,
             "image_path": None, "icons": []},
        ],
    }
    seed = PPTGenerationState(topic="测试主题", requirements="初始")
    seed.set("outline", outline)
    app_module._save_state(seed)

    resp = client.post("/generate", data={
        "topic": "测试主题",
        "requirements": "全部改成学术风",
        "filename": "global_route.pptx",
        "mode": "modify_text",
        "modify_page": "",
    })
    assert resp.status_code == 200
    body = resp.get_data(as_text=True)
    assert "全局仅改文本" in body, "页面应显示全局修改标签"
    print("PASS: 全局修改路由（页面标签含「全局仅改文本」）\n")


if __name__ == "__main__":
    test_parse_page_numbers()
    main()
    test_global_routing()
