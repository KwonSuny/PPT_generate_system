import os
import json
import re
import traceback
from tempfile import mkstemp
from copy import deepcopy
from datetime import datetime
from pathlib import Path
from typing import List, Optional

from agents.planner import get_planner_agent
from agents.writer import get_writer_agent
from agents.review import review_ppt_outline
from agents.image_generator import get_image_generator 
from agents.template import get_layout_agent
from renderer import PPTXRenderer
from config.schema import PPTOutline, PPTFinalBlueprint, PagePolicy
from config.content import (ContentConstraintError, parse_page_policy, check_page_count, identified_outline,
                            identified_slide, lock_slide, bind_canvas, lock_text_canvas)
from config.debug import save_agent_log 
from config.state import PPTGenerationState
from config.llm_client import resolve_model_name
from config.sanitizers import snap_font_size
from pydantic import ValidationError


def _export_pptx(filename, export, expected_slides=None):
    from pptx import Presentation
    output = Path(filename).resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary = mkstemp(prefix=f".{output.stem}_", suffix=".pptx", dir=output.parent)
    os.close(descriptor)
    try:
        export(temporary)
        actual_count = len(Presentation(temporary).slides)
        if not actual_count:
            raise ValueError("导出文件不含任何页面")
        if expected_slides is not None and actual_count != expected_slides:
            raise ContentConstraintError("导出文件页数与已确认的蓝图不一致")
        with open(temporary, "rb+") as stream:
            os.fsync(stream.fileno())
        os.replace(temporary, output)
    finally:
        Path(temporary).unlink(missing_ok=True)
    return str(output)


def _render_blueprint(blueprint, filename):
    return _export_pptx(filename, PPTXRenderer(blueprint).render_all, len(blueprint.canvases))


def _layout_page(state, designer, payload, page_number, run_log_dir, stage_prefix="layout", canvas_validator=None, canvas_transform=None):
    """已知排版约束错误最多反馈两轮，原文与字号下限保持不变。"""
    from config.layout_budget import capacity_budget, layout_issues, layout_fingerprint, font_floor, design_request_policy, font_resize_locked
    policy = payload.get('layout_policy') or design_request_policy(state.requirements)
    resize = policy.get('allow_font_resize') or (stage_prefix=='layout' and (state.operation_mode!='incremental_page' or state.modify_scope=='design')
        and not font_resize_locked(state.requirements))
    policy = {**policy, 'allow_font_resize': bool(resize)}
    payload = {**payload, 'layout_policy':policy}
    floors = dict((payload.get('layout_feedback') or {}).get('text_font_size_floors', {}))
    if payload['slides'][0].get('table'):
        floors.pop('table/source',None)
    component_floors = dict((payload.get('layout_feedback') or {}).get('component_font_size_floors', {}))
    seeded_floors = bool(floors or component_floors)
    budget = capacity_budget(payload['slides'][0], floors, component_floors)
    request = {**deepcopy(payload), 'capacity_budget': budget}
    save_agent_log(budget, f'layout_page_{page_number}_budget.json', run_log_dir)
    last_error = None
    rejected = list((payload.get('layout_feedback') or {}).get('rejected_layouts', []))
    captured = False
    for attempt in range(3):
        key = f"{stage_prefix}/page{page_number}" + (f"/repair{attempt}" if attempt else "")
        blueprint = None
        bound = None
        resolved = None
        try:
            blueprint = state.execute_stage(
                key, request, lambda: designer(state, request),
                model=resolve_model_name("layout"), result_model=PPTFinalBlueprint,
            )
            save_agent_log(blueprint.model_dump(), f"layout_page_{page_number}_attempt_{attempt + 1}.json", run_log_dir)
            if len(blueprint.canvases) != 1:
                raise ContentConstraintError(f"Layout 必须为第 {page_number} 页返回严格一个画布")
            bound = bind_canvas(blueprint.canvases[0], payload["slides"][0], preserve_component_style=policy.get('allow_theme_change',False))
            bound.layout_resolved = False
            native_source = next((e.table.source for e in bound.elements if e.type=='table' and e.table and e.table.source),None)
            if native_source:
                # 来源已由原生表格绘制，避免另一个来源框重复占据正文空间。
                for element in bound.elements:
                    if element.type=='textbox':
                        element.paragraphs=[p for p in element.paragraphs or [] if p.content_ref!='table/source']
                bound.elements=[e for e in bound.elements if e.type!='textbox' or e.paragraphs]
            shown = {p.content_ref for e in bound.elements for p in e.paragraphs or []}
            if native_source:
                shown.add('table/source')
            missing = set(payload.get('required_content_refs',[]))-shown
            if missing:
                raise ContentConstraintError('排版遗漏原版文字引用：'+'、'.join(sorted(missing)))
            for element in bound.elements:
                if element.type=='textbox':
                    for paragraph in element.paragraphs or []:
                        if (paragraph.content_ref or '').startswith('full_content/'):
                            paragraph.font_size = max(font_floor(paragraph.content_ref,state.requirements), paragraph.font_size)
                elif element.type=='icon' and not element.content_ref:
                    nearby = [tb for tb in bound.elements if tb.type=='textbox' and len(tb.paragraphs or [])==1
                              and 0 <= tb.left-element.left-element.width <= .6 and abs(tb.top-element.top) <= .5]
                    if len(nearby)==1:
                        element.content_ref = nearby[0].paragraphs[0].content_ref
            if canvas_transform:
                bound = canvas_transform(bound)
            if not captured:
                for element in bound.elements:
                    for paragraph in (element.paragraphs or []) if element.type == "textbox" else []:
                        if paragraph.content_ref:
                            minimum = font_floor(paragraph.content_ref,state.requirements) if resize else snap_font_size(paragraph.font_size)
                            floors[paragraph.content_ref] = max(floors.get(paragraph.content_ref, 0), minimum)
                    if element.type in {"table", "diagram"}:
                        style = getattr(element, element.type + "_style")
                        previous = component_floors.get(element.type, {})
                        component_floors[element.type] = {name: max((12 if name=='source_font_size' else 14 if name=='label_font_size' else 18) if resize else value, previous.get(name, 0))
                            for name, value in style.model_dump().items() if "font_size" in name}
                captured = True
            try:
                resolved = PPTXRenderer.resolve_canvas(bound, fit_text=True, reflow_text=attempt == 2)
            except ContentConstraintError:
                if not resize:
                    raise
                fitted = PPTXRenderer.fit_readable_text(bound, floors)
                resolved = PPTXRenderer.resolve_canvas(fitted, fit_text=True, reflow_text=attempt == 2)
            if captured or seeded_floors:
                covered = {paragraph.content_ref for element in resolved.elements if element.type == "textbox" for paragraph in element.paragraphs or []}
                if floors.keys() - covered:
                    raise ContentConstraintError("排版修正删除了已展示的文字引用：" + "、".join(sorted(floors.keys() - covered)))
                for element in resolved.elements:
                    for paragraph in (element.paragraphs or []) if element.type == "textbox" else []:
                        if paragraph.font_size < floors.get(paragraph.content_ref, 0):
                            raise ContentConstraintError("排版修正缩小了字号，必须调整容器而不是缩字")
                    if element.type in component_floors:
                        style = getattr(element, element.type + "_style")
                        if any(getattr(style, name) < minimum for name, minimum in component_floors[element.type].items()):
                            raise ContentConstraintError("排版修正缩小了结构化图元字号")
            if canvas_validator:
                canvas_validator(resolved)
            issues = layout_issues(resolved)
            if issues:
                raise ContentConstraintError('排版检查未通过：'+'；'.join(
                    item['message']+'('+','.join(str(v) for v in item['element_ids'])+')' for item in issues))
            if stage_prefix == "layout":
                state.set("layout_recovery_request", None)
                state.set("layout_fallbacks", [item for item in state.get("layout_fallbacks", []) if item["page"] != page_number])
            save_agent_log(resolved.model_dump(), f"layout_page_{page_number}_resolved.json", run_log_dir)
            return resolved
        except (ContentConstraintError, ValidationError) as exc:
            last_error = str(exc)
            feedback = {"error": last_error, "text_font_size_floors": floors,
                        "component_font_size_floors": component_floors,
                        "instruction": "仅重新排版本次页面；保留内容引用、数据、SVG、图片和页数。先调整分区，允许缩放时可降低至明确的字号下限，不能删改内容。",
                        'layout_policy':policy}
            if blueprint and len(blueprint.canvases) == 1:
                feedback_canvas = bound if bound is not None else blueprint.canvases[0]
                feedback["previous_canvas"] = feedback_canvas.model_dump(exclude_none=True, exclude_defaults=True)
                feedback["required_text_heights"] = [
                    {"id": element.id, "minimum_height": round(PPTXRenderer.required_text_height(element) + 0.01, 3)}
                    for element in feedback_canvas.elements if element.type == "textbox"
                ]
                diagnostic_canvas = resolved if resolved is not None else feedback_canvas
                # resolve 的确定性对齐也要进入反馈，不能只让模型看到处理前的猜测。
                feedback['previous_canvas'] = diagnostic_canvas.model_dump(exclude_none=True, exclude_defaults=True)
                feedback['violations'] = layout_issues(diagnostic_canvas)
                fingerprint = layout_fingerprint(diagnostic_canvas)
                feedback['repeated_geometry'] = fingerprint in rejected
                rejected.append(fingerprint)
                feedback['rejected_layouts'] = rejected.copy()
                feedback['repair_plan'] = ('上次仍使用相同分区，必须切换阅读轴或图文构图。' if feedback['repeated_geometry'] else '') + \
                    '先锁定文字和图示所需空间，再为每列预留图标与邻区间距；按 violations 中的坐标和差额移动关联分组。图示不能自动缩小避让。'
            budget = capacity_budget(payload['slides'][0], floors, component_floors)
            feedback['capacity_budget'] = budget
            save_agent_log(feedback, f"layout_page_{page_number}_error_{attempt + 1}.json", run_log_dir)
            if attempt < 2:
                state.append_progress(f"Layout: 第 {page_number} 页约束未通过，反馈修正 {attempt + 1}/2：{last_error}")
                request = {**deepcopy(payload), 'capacity_budget': budget, "layout_feedback": feedback}
    if stage_prefix == "layout" and not canvas_validator and not canvas_transform:
        from config.layout_fallback import stable_layout, StableLayoutUnavailable
        from config.schema import SlideCanvas
        def fallback():
            try:
                canvas, composition = stable_layout(payload["slides"][0], floors, component_floors)
            except StableLayoutUnavailable:
                if not resize:
                    raise
                canvas, composition = stable_layout(payload["slides"][0], floors, component_floors, compact=True)
            records = [item for item in state.get("layout_fallbacks", []) if item["page"] != page_number]
            records.append({"page": page_number, "slide_id": canvas.slide_id, "reason": last_error,
                            "composition": composition,
                            "visual_warnings": [item for item in layout_issues(canvas) if item['kind']=='svg_readability']})
            state.set("layout_fallbacks", records)
            state.set("layout_recovery_request", None)
            state.append_progress(f"Layout: 第 {page_number} 页自由排版未通过，已切换保留完整内容的稳定布局")
            return canvas
        try:
            resolved = state.execute_stage(f"{stage_prefix}/page{page_number}/fallback",
                {"slide": payload["slides"][0], "text_floors": floors, "component_floors": component_floors,
                 "reason": last_error}, fallback, external=False, result_model=SlideCanvas)
        except StableLayoutUnavailable as exc:
            state.set("layout_recovery_request", {"page": page_number, "message": str(exc),
                      "capacity_conflict": exc.capacity_conflict})
            raise
        save_agent_log(resolved.model_dump(), f"layout_page_{page_number}_resolved.json", run_log_dir)
        save_agent_log(state.get("layout_fallbacks", []), f"layout_page_{page_number}_fallback.json", run_log_dir)
        return resolved
    raise ContentConstraintError(f"第 {page_number} 页排版在两轮修正后仍不满足约束：{last_error}；已保留内容与每轮蓝图，未覆盖成品")


def _render_offline_demo(user_topic: str, filename: str, page_policy=None) -> str:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt

    output_path = os.path.abspath(filename)
    output_dir = os.path.dirname(output_path)
    if output_dir:
        os.makedirs(output_dir, exist_ok=True)

    lines = [line.strip() for line in user_topic.splitlines() if line.strip()]
    title = lines[0] if lines else "PPT 生成演示"
    subtitle = "本地离线演示模式" if len(lines) <= 1 else lines[1]
    bullets = []
    for line in lines[2:]:
        if line.startswith("要求") or line.startswith("提示"):
            continue
        bullets.append(line)
    if not bullets:
        bullets = ["主题已接收", "可继续扩展内容", "适配网页生成流程"]

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Cover slide
    cover = prs.slides.add_slide(blank_layout)
    cover.background.fill.solid()
    cover.background.fill.fore_color.rgb = RGBColor(0xF8, 0xFA, 0xFC)
    title_box = cover.shapes.add_textbox(Inches(0.8), Inches(1.2), Inches(11.5), Inches(1.2))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(26)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0x11, 0x1F, 0x2E)
    p.alignment = PP_ALIGN.LEFT

    subtitle_box = cover.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(10.5), Inches(1.0))
    tf2 = subtitle_box.text_frame
    tf2.word_wrap = True
    p2 = tf2.paragraphs[0]
    p2.text = subtitle
    p2.font.size = Pt(16)
    p2.font.color.rgb = RGBColor(0x65, 0x6B, 0x7A)

    content = prs.slides.add_slide(blank_layout)
    content.background.fill.solid()
    content.background.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    box = content.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.5), Inches(5.6))
    tf3 = box.text_frame
    tf3.word_wrap = True
    tf3.margin_left = Inches(0.08)
    tf3.margin_right = Inches(0.08)
    title_p = tf3.paragraphs[0]
    title_p.text = "生成内容摘要"
    title_p.font.size = Pt(20)
    title_p.font.bold = True
    title_p.font.color.rgb = RGBColor(0x0F, 0x5B, 0xA4)

    for bullet in bullets[:5]:
        p = tf3.add_paragraph()
        p.text = f"• {bullet}"
        p.font.size = Pt(16)
        p.font.color.rgb = RGBColor(0x33, 0x3F, 0x4B)
        p.level = 0
        p.space_after = Pt(8)

    summary = prs.slides.add_slide(blank_layout)
    summary.background.fill.solid()
    summary.background.fill.fore_color.rgb = RGBColor(0xF8, 0xFA, 0xFC)
    box2 = summary.shapes.add_textbox(Inches(0.8), Inches(1.2), Inches(11.2), Inches(4.2))
    tf4 = box2.text_frame
    tf4.word_wrap = True
    p4 = tf4.paragraphs[0]
    p4.text = "网页端已成功接入 PPT 生成流程"
    p4.font.size = Pt(22)
    p4.font.bold = True
    p4.font.color.rgb = RGBColor(0x0F, 0x5B, 0xA4)
    p5 = tf4.add_paragraph()
    p5.text = "如配置了 DASHSCOPE_API_KEY，系统会继续走完整的 Planner / Writer / Layout 智能体链路。"
    p5.font.size = Pt(16)
    p5.font.color.rgb = RGBColor(0x65, 0x6B, 0x7A)

    policy = PagePolicy.model_validate(page_policy or {})
    count = max(policy.min_slides, min(3, policy.max_slides))
    while len(prs.slides) > count:
        slide_id = prs.slides._sldIdLst[-1]
        prs.part.drop_rel(slide_id.rId)
        prs.slides._sldIdLst.remove(slide_id)
    while len(prs.slides) < count:
        slide = prs.slides.add_slide(blank_layout)
        text = slide.shapes.add_textbox(Inches(0.8), Inches(1.2), Inches(11.5), Inches(2)).text_frame
        text.paragraphs[0].text = f"离线演示第 {len(prs.slides)} 页：{title}"
        text.paragraphs[0].font.size = Pt(26)
    return _export_pptx(output_path, prs.save, count)


def run_incremental_modify(
    user_topic: str,
    page_number: int,
    filename: str = "output.pptx",
    state: Optional[PPTGenerationState] = None,
    requirements: str = "",
    run_dir: Optional[str] = None,
):
    """单页增量修改（兼容入口）：转调多页增量管线。"""
    return run_incremental_modify_pages(
        user_topic,
        page_numbers=[page_number],
        filename=filename,
        state=state,
        requirements=requirements,
        run_dir=run_dir,
    )


def run_incremental_modify_pages(
    user_topic: str,
    page_numbers: Optional[List[int]],
    filename: str = "output.pptx",
    state: Optional[PPTGenerationState] = None,
    requirements: str = "",
    run_dir: Optional[str] = None,
):
    """
    页面级增量修改管线（支持多页 / 全局）：仅对指定页集合调用 Writer / 生图 / Layout，
    其余页面完全沿用缓存的大纲与物理蓝图，最后整体重渲染导出。

    :param user_topic: 用户本轮输入（主题 + 最新修改要求）
    :param page_numbers: 目标页码列表（从 1 开始）；为空或 None 表示全局（全部页）
    :param filename: 最终导出的 pptx 文件路径
    :param state: 携带缓存 outline/blueprint 的会话状态
    :param requirements: 用户最新要求
    :param run_dir: 统一的大时间戳文件夹路径；为空时回退 debug_logs/run_{ts}
    """
    if state is None:
        state = PPTGenerationState(topic=user_topic, requirements=requirements)
    else:
        state.topic = user_topic
        state.requirements = requirements

    state.set("generation_status", "running")
    state.set("result_kind", "generated")
    state.set("content_review_status", "not_run")
    cached_outline = deepcopy(state.get("outline"))
    cached_blueprint = deepcopy(state.get("blueprint"))
    slides = cached_outline.get("slides") if isinstance(cached_outline, dict) else None
    canvases = cached_blueprint.get("canvases") if isinstance(cached_blueprint, dict) else None

    def _fail(reason: str) -> Optional[str]:
        print(f"[Incremental] {reason}，已停止页面修改。")
        # 恢复缓存，避免中途失败污染历史状态
        if isinstance(cached_outline, dict) and cached_outline:
            state.set("outline", cached_outline)
        if isinstance(cached_blueprint, dict) and cached_blueprint:
            state.set("blueprint", cached_blueprint)
        state.set("generation_status", "failed")
        state.append_progress(f"Error: {reason}，已保留原缓存")
        return None

    api_key = (os.environ.get("DASHSCOPE_API_KEY") or "").strip()
    if not api_key:
        return _fail("未配置模型密钥，无法执行页面修改")

    if not isinstance(slides, list) or not slides or not isinstance(canvases, list) or not canvases:
        return _fail("缓存的大纲或物理蓝图缺失/非法")
    total = len(slides)
    if total != len(canvases):
        return _fail(f"大纲与蓝图页数不对齐（{total} vs {len(canvases)}）")

    # 解析目标页集合：None/空 → 全局（全部页）
    if page_numbers is not None and not isinstance(page_numbers, (list, tuple)):
        return _fail("目标页码必须是整数列表")
    targets = list(page_numbers) if page_numbers else list(range(1, total + 1))
    if any(type(p) is not int for p in targets):
        return _fail("目标页码必须是整数")
    targets = sorted(set(targets))
    if not targets:
        return _fail("目标页集合为空")
    if any(p < 1 or p > total for p in targets):
        return _fail(f"页码 {targets} 超出范围（当前共 {total} 页）")

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    run_log_dir = os.path.join(run_dir, "run") if run_dir else f"debug_logs/run_{timestamp}"
    topic = cached_outline.get("topic") or state.topic

    state.append_progress(f"Incremental: 开始 {', '.join(map(str, targets))} 页增量修改")
    print(f"\n[Incremental] 目标页 {targets}（共 {total} 页）增量修改启动，其余页面沿用缓存。")

    try:
        writer = get_writer_agent(api_key=api_key, model_name=resolve_model_name("writer"))
        if state.modify_scope != "text":
            image_generator = get_image_generator(api_key=api_key, model_name=resolve_model_name("image"))
            layout_designer = get_layout_agent(api_key=api_key, model_name=resolve_model_name("layout"))

        merged_slides = list(slides)
        merged_canvases = list(canvases)

        for page_number in targets:
            page_idx = page_number - 1
            target_slide = identified_slide(slides[page_idx], page_number)

            # 邻页主题色：注入输入 payload，约束单页修改的色彩家族连续性
            neighbor_themes = [
                slides[i].get("theme")
                for i in (page_idx - 1, page_idx + 1)
                if 0 <= i < total and isinstance(slides[i], dict)
            ]
            neighbor_themes = [t for t in neighbor_themes if isinstance(t, dict)]

            # ---- 文本/主题层：仅目标页改写（design 也走 Writer，但锁文字只改 theme） ----
            single_payload = {"schema_version": 2, "topic": topic, "slides": [target_slide]}
            if neighbor_themes:
                single_payload["neighbor_themes"] = neighbor_themes
            updated_outline = state.execute_stage(
                f"writer/page{page_number}", single_payload, lambda: writer(state, single_payload),
                model=resolve_model_name("writer"), result_model=PPTOutline,
            )
            if len(updated_outline.slides) != 1:
                raise ContentConstraintError("Writer 必须为目标页返回严格一页；已停止修改")
            updated_slide = identified_slide(lock_slide(target_slide, updated_outline.slides[0].model_dump(), state.modify_scope), page_number)
            if state.modify_scope == "text":
                new_canvas = lock_text_canvas(canvases[page_idx], target_slide, updated_slide)
                merged_slides[page_idx] = updated_slide
                merged_canvases[page_idx] = new_canvas.model_dump()
                state.append_progress(f"Incremental: 第 {page_number} 页文字已回填，保留原物理布局")
                continue

            # ---- 图像层：仅目标页生图（SVG 页跳过，与全量管线策略一致） ----
            has_component = updated_slide.get("layout_type") in {"svg", "table", "diagram"} or updated_slide.get("table") or updated_slide.get("diagram") or (updated_slide.get("svg_code") or "").strip()
            if has_component:
                updated_slide["image_path"] = None
                state.append_progress(f"Incremental: 第 {page_number} 页使用结构化图示，跳过生图")
            elif (updated_slide.get("image_prompt") or "").strip():
                state.append_progress(f"Incremental: 仅生成第 {page_number} 页插图")
                theme = updated_slide.get("theme") or {}
                enhanced_prompt = f"""
                    {updated_slide.get('image_prompt')}。画面主色调采用：{theme.get('primary')}，辅助高亮色采用：{theme.get('secondary')}。
                    这是一张 16:9 演示文稿辅助插图，只用于侧边、角落或独立插图区点缀，不作为背景图或文字衬底。
                    画面需要主体清晰、边缘可裁切、构图简洁，避免中心堆满复杂纹理。
                    严禁生成任何具体的英文字母、汉字或文本，严禁生成复杂统计指标、柱状图等需要精确阅读的图表。
                    避免高对比杂乱纹理、暗部噪点和会干扰正文阅读的复杂背景；整体清爽、克制。
                    """
                local_img_path = state.execute_stage(
                    f"image/page{page_number}", {"prompt": enhanced_prompt, "size": "1280x720"},
                    lambda: image_generator(prompt=enhanced_prompt, size="1280x720", output_dir=run_log_dir,
                                            filename=f"slide_accent_{page_idx}.png"),
                    model=resolve_model_name("image"),
                )
                if not local_img_path or not Path(local_img_path).is_file():
                    raise RuntimeError(f"第 {page_number} 页配图生成失败")
                updated_slide["image_path"] = local_img_path
            else:
                updated_slide["image_path"] = None

            # ---- 物理层：仅目标页重新排版 ----
            state.append_progress(f"Incremental: 生成第 {page_number} 页物理蓝图")
            layout_payload = {"schema_version": 2, "topic": topic, "slides": [updated_slide]}
            if neighbor_themes:
                layout_payload["neighbor_themes"] = neighbor_themes
            new_canvas = _layout_page(state, layout_designer, layout_payload, page_number, run_log_dir)

            # ---- 拼接目标页，其余页保持缓存原样 ----
            merged_slides[page_idx] = updated_slide
            merged_canvases[page_idx] = new_canvas.model_dump()

        merged_outline = {**cached_outline, "schema_version": 2, "topic": topic, "slides": merged_slides}
        merged_payload = {**cached_blueprint, "schema_version": 2, "canvases": merged_canvases}
        merged_blueprint = PPTFinalBlueprint.model_validate(merged_payload)
        state.set("outline", merged_outline)
        state.set("blueprint", merged_payload)
        save_agent_log(merged_outline, "incremental_outline.json", run_log_dir)
        save_agent_log(merged_payload, "incremental_blueprint.json", run_log_dir)

        state.execute_stage("render", {"blueprint": merged_blueprint.model_dump()},
                            lambda: _render_blueprint(merged_blueprint, filename), external=False)
        state.set("generation_status", "succeeded")
        state.append_progress("Incremental: 完成，PPT 导出成功")
        print(f" [Incremental] 目标页 {targets} 增量修改闭环，成功导出至: {filename}")
        return filename
    except Exception as e:
        print(f"\n[Incremental] 增量修改发生崩溃！故障报告: {e}")
        traceback_text = traceback.format_exc()
        os.makedirs(run_log_dir, exist_ok=True)
        save_agent_log({"error": str(e), "pages": targets}, "error.json", run_log_dir)
        error_txt_path = os.path.join(run_log_dir, "error.txt")
        with open(error_txt_path, "w", encoding="utf-8") as f:
            f.write(traceback_text)
        return _fail(str(e) if isinstance(e, ContentConstraintError) else "增量修改过程发生异常，详情见本次运行日志")


def run_ppt_system(
    user_topic: str, 
    filename: str = "output.pptx",
    from_content_json: Optional[str] = None,
    from_blueprint_json: Optional[str] = None,
    state: Optional[PPTGenerationState] = None,
    requirements: str = "",
    run_dir: Optional[str] = None,
    page_range: Optional[str] = None,
):
    """
    PPT 自动化高阶管线总线
    
    :param user_topic: PPT 主题描述
    :param filename: 最终导出的 pptx 文件路径
    :param from_content_json: 可选，本地缓存的 PPTOutline JSON 路径。传入后将跳过大纲生产和生图阶段。
    :param from_blueprint_json: 可选，本地缓存的 PPTFinalBlueprint JSON 路径。传入后将直接渲染。
    :param run_dir: 统一的大时间戳文件夹路径；为空时回退 debug_logs/run_{ts}
    """
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    run_log_dir = os.path.join(run_dir, "run") if run_dir else f"debug_logs/run_{timestamp}"
    if state is None:
        state = PPTGenerationState(topic=user_topic, requirements=requirements)
    else:
        state.topic = user_topic
        state.requirements = requirements
    state.set("generation_status", "running")
    state.set("result_kind", "generated")
    state.set("content_review_status", "not_run")

    if from_blueprint_json:
        print(f"\n 检测到物理蓝图缓存注入：{from_blueprint_json}")
        if not os.path.exists(from_blueprint_json):
            print(f"错误：指定的蓝图缓存文件不存在！")
            state.set("generation_status", "failed")
            state.append_progress("Error: 指定的蓝图缓存文件不存在")
            return None
        
        try:
            with open(from_blueprint_json, "r", encoding="utf-8") as f:
                blueprint_dict = json.load(f)
            
            # 反序列化为 Pydantic 强校验对象
            final_blueprint = PPTFinalBlueprint.model_validate(blueprint_dict)
            if not final_blueprint.canvases:
                raise ValueError("物理蓝图不含任何页面")
            print("物理蓝图反序列化成功！")
            
            # 直接移交哑巴渲染器
            state.execute_stage("render", {"blueprint": final_blueprint.model_dump()},
                                lambda: _render_blueprint(final_blueprint, filename), external=False)
            state.set("blueprint", final_blueprint.model_dump())
            state.set("generation_status", "succeeded")
            state.append_progress("Rendering: 蓝图缓存回放完成")
            return filename
        except Exception as e:
            print(f"蓝图注入并渲染失败: {e}")
            state.set("generation_status", "failed")
            state.append_progress("Error: 蓝图缓存注入或渲染失败")
            return None

    api_key = (os.environ.get("DASHSCOPE_API_KEY") or "").strip()
    try:
        policy = parse_page_policy(page_range, user_topic + "\n" + requirements) if page_range is not None or not state.get("page_policy") else PagePolicy.model_validate(state.get("page_policy"))
        state.set("page_policy", policy.model_dump())
    except (ValueError, ContentConstraintError) as exc:
        state.set("generation_status", "failed")
        state.append_progress("Error: " + (str(exc) if isinstance(exc, ContentConstraintError) else "页数约束无效"))
        return None
    if not api_key:
        if from_content_json:
            state.set("generation_status", "failed")
            state.append_progress("Error: 内容缓存仍需调用排版模型，请配置模型密钥")
            return None
        print("未检测到 DASHSCOPE_API_KEY，使用离线演示模式生成基础 PPT。")
        state.set("result_kind", "offline_demo")
        try:
            demo_path = state.execute_stage("demo", {"topic": user_topic, "page_policy": policy.model_dump()},
                                            lambda: _render_offline_demo(user_topic, filename, policy), external=False)
        except Exception:
            state.set("generation_status", "failed")
            state.append_progress("Error: 离线演示文件导出失败")
            return None
        state.set("generation_status", "succeeded")
        state.append_progress("Demo: 离线演示文件已导出，未执行智能体生成与审核")
        return demo_path

    if not state.messages or state.messages[-1].get("role") != "user":
        state.append_message("user", user_topic if not requirements else f"{user_topic}\n要求：{requirements}")
    final_content = None 

    if from_content_json:
        print(f"\n检测到文本语义大纲缓存注入：{from_content_json}")
        if not os.path.exists(from_content_json):
            print(f"错误：指定的内容缓存文件不存在！")
            state.set("generation_status", "failed")
            state.append_progress("Error: 指定的内容缓存文件不存在")
            return None
        
        try:
            with open(from_content_json, "r", encoding="utf-8") as f:
                content_dict = json.load(f)
            final_content = PPTOutline.model_validate(content_dict)
            print("语义大纲反序列化成功！")
        except Exception as e:
            print(f"内容大纲注入失败: {e}")
            state.set("generation_status", "failed")
            state.append_progress("Error: 内容缓存注入失败")
            return None
            
    else:
        try:
            planner = get_planner_agent(api_key=api_key, model_name=resolve_model_name("planner"))
            writer = get_writer_agent(api_key=api_key, model_name=resolve_model_name("writer"))
            image_generator = get_image_generator(api_key=api_key, model_name=resolve_model_name("image"))
            current_feedback = None
            max_retries = 3
            attempt = 0
            compiled_content = None
            
            # 内容生产与质检闭环
            while attempt < max_retries:
                attempt += 1
                progress_msg = f"第 {attempt} 轮智能体协同" 
                state.append_progress(progress_msg)
                print(f"\n {progress_msg} ")
                
                raw_outline = state.execute_stage(
                    f"planner/attempt{attempt}", {"feedback": current_feedback},
                    lambda: planner(state, feedback_list=current_feedback),
                    model=resolve_model_name("planner"), result_model=PPTOutline,
                )
                save_agent_log(raw_outline.model_dump(), f"raw_outline_attempt_{attempt}.json", run_log_dir)
                try:
                    check_page_count(len(raw_outline.slides), policy)
                except ContentConstraintError as exc:
                    current_feedback = [str(exc)]
                    continue
                compiled_content = state.execute_stage(
                    f"writer/attempt{attempt}", raw_outline.model_dump(), lambda: writer(state, raw_outline),
                    model=resolve_model_name("writer"), result_model=PPTOutline,
                )
                save_agent_log(compiled_content.model_dump(), f"compiled_content_attempt_{attempt}.json", run_log_dir)
                
                # 检查内容合规性
                check_page_count(len(compiled_content.slides), policy)
                is_passed, errors = review_ppt_outline(compiled_content, page_policy=policy)
                
                if is_passed:
                    state.set("content_review_status", "passed")
                    print(f"[Reviewer] 内容终审通过！准备唤醒生图和物理画布排版。")
                    final_content = compiled_content
                    save_agent_log(final_content.model_dump(), f"final_content_passed.json", run_log_dir)
                    break
                else:
                    print(f"[Reviewer] 质检失败！已派发回上一级。")
                    current_feedback = errors  
                    
            if not final_content:
                if compiled_content is None:
                    raise ContentConstraintError(current_feedback[0] if current_feedback else "未获得符合页数要求的内容")
                print(f"\n系统提示：未能在限制轮次内完美对齐，采用最后一轮输出保底。")
                final_content = compiled_content
                state.set("content_review_status", "needs_review")
                state.append_progress("Review: 内容检查未通过，导出结果需要人工复核")

            # 调用图像模型生成图并写入路径
            print("\n正在遍历各幻灯片，触发图片生成管道...")
            for i, slide in enumerate(final_content.slides):
                if slide.layout_type in {"svg", "table", "diagram"} or slide.table or slide.diagram or (slide.svg_code and slide.svg_code.strip()):
                    slide.image_path = None
                    state.append_progress(f"Images: 第 {i + 1} 页使用表格、关系图或 SVG，跳过生图")
                    print(f"[Pipeline] 第 {i + 1} 页使用表格、关系图或 SVG，已跳过生图。")
                    continue
                    
                if slide.image_prompt and slide.image_prompt.strip():
                    state.append_progress(f"Images: 生成第 {i + 1} 页插图")
                    enhanced_prompt = f"""
                    {slide.image_prompt}。画面主色调采用：{slide.theme.primary}，辅助高亮色采用：{slide.theme.secondary}。
                    这是一张 16:9 演示文稿辅助插图，只用于侧边、角落或独立插图区点缀，不作为背景图或文字衬底。
                    画面需要主体清晰、边缘可裁切、构图简洁，避免中心堆满复杂纹理。
                    严禁生成任何具体的英文字母、汉字或文本，严禁生成复杂统计指标、柱状图等需要精确阅读的图表。
                    避免高对比杂乱纹理、暗部噪点和会干扰正文阅读的复杂背景；整体清爽、克制。
                    """
                    local_img_path = state.execute_stage(
                        f"image/page{i + 1}", {"prompt": enhanced_prompt, "size": "1280x720"},
                        lambda: image_generator(prompt=enhanced_prompt, size="1280x720", output_dir=run_log_dir,
                                                filename=f"slide_accent_{i}.png"),
                        model=resolve_model_name("image"),
                    )
                    if not local_img_path or not Path(local_img_path).is_file():
                        raise RuntimeError(f"第 {i + 1} 页配图生成失败")
                    slide.image_path = local_img_path
            
            save_agent_log(final_content.model_dump(), f"final_content_with_images.json", run_log_dir)
            state.append_progress("Images: 生成并写入插图路径")
        except Exception as e:
            print(f"\n内容生成或图像生成阶段发生崩溃！故障报告: {e}")
            traceback_text = traceback.format_exc()
            os.makedirs(run_log_dir, exist_ok=True)
            save_agent_log({"error": str(e)}, "error.json", run_log_dir)
            error_txt_path = os.path.join(run_log_dir, "error.txt")
            with open(error_txt_path, "w", encoding="utf-8") as f:
                f.write(traceback_text)
            state.set("generation_status", "failed")
            state.append_progress("Error: " + (str(e) if isinstance(e, ContentConstraintError) else "内容或图片生成失败，详情见本次运行日志"))
            return None

    # 视觉排版设计与物理画布渲染管线
    # 逐页排版：每次仅向 Layout Agent 提交单页语义大纲，避免一次性生成全部页物理蓝图，
    # 导致单次请求输出体积过大 / 生成耗时过长，触发 MaaS 网关主动断连（Connection error）。
    try:
        final_content = identified_outline(final_content.model_dump())
        state.set("outline", final_content.model_dump())
        layout_designer = get_layout_agent(api_key=api_key, model_name=resolve_model_name("layout"))
        writer_dump_data = final_content.model_dump()
        slides = writer_dump_data.get("slides") or []
        total = len(slides)
        if not total:
            raise ValueError("语义大纲不含任何页面，无法排版")
        topic = writer_dump_data.get("topic") or state.topic

        state.append_progress(f"Rendering: 开始逐页布局渲染（共 {total} 页）")
        merged_canvases = []
        state.set("layout_fallbacks", [])
        state.set("layout_recovery_request", None)
        for i, slide in enumerate(slides):
            state.append_progress(f"Rendering: 排版第 {i + 1}/{total} 页")
            print(f" [Layout] 逐页排版第 {i + 1}/{total} 页")
            single_payload = {"schema_version": 2, "topic": topic, "slides": [slide]}
            # 邻页主题色：约束单页配色的色彩家族连续性（与增量管线一致）
            neighbor_themes = [
                slides[j].get("theme")
                for j in (i - 1, i + 1)
                if 0 <= j < total and isinstance(slides[j], dict)
            ]
            neighbor_themes = [t for t in neighbor_themes if isinstance(t, dict)]
            if neighbor_themes:
                single_payload["neighbor_themes"] = neighbor_themes

            merged_canvases.append(_layout_page(state, layout_designer, single_payload, i + 1, run_log_dir).model_dump())

        final_blueprint = PPTFinalBlueprint.model_validate({"schema_version": 2, "canvases": merged_canvases})
        # 显式写回全量蓝图，确保落盘快照（_persist_generation_assets 依赖 state["blueprint"]）
        state.set("blueprint", final_blueprint.model_dump())
        save_agent_log(final_blueprint.model_dump(), f"final_blueprint.json", run_log_dir)
        state.execute_stage("render", {"blueprint": final_blueprint.model_dump()},
                            lambda: _render_blueprint(final_blueprint, filename), external=False)
        state.set("generation_status", "succeeded")
        state.append_progress("Rendering: 完成，PPT 导出成功")
        print(f" 幻灯片系统闭环，成功导出至: {filename}")
    except Exception as e:
        print(f" 视觉排版或物理渲染阶段发生崩溃！故障报告: {e}")
        traceback_text = traceback.format_exc()
        os.makedirs(run_log_dir, exist_ok=True)
        save_agent_log({"error": str(e)}, "error.json", run_log_dir)
        error_txt_path = os.path.join(run_log_dir, "error.txt")
        with open(error_txt_path, "w", encoding="utf-8") as f:
            f.write(traceback_text)
        state.set("generation_status", "failed")
        state.append_progress("Error: " + (str(e) if isinstance(e, ContentConstraintError) else "排版或 PPTX 导出失败，详情见本次运行日志"))
        return None

    return filename


if __name__ == "__main__":
    raise SystemExit("请使用 python app.py 与 python worker.py，或 python launcher.py 启动工作台。")
