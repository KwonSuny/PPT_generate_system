import json
import os
import re
from datetime import datetime
from functools import lru_cache
from pathlib import Path
from tempfile import mkdtemp
from typing import Optional
from uuid import uuid4

from flask import Flask, jsonify, render_template, request, send_file, session
from pptx import Presentation

from config.jobs import dispatch_pending
from config.content import ContentConstraintError, parse_page_policy, check_page_count
from config.state import PPTGenerationState
from config.storage import LOCAL_OWNER, PPTStore, StorageConflict, content_hash
from config_manager import (
    clear_api_key, get_output_root, has_api_key, load_base_url, save_api_key, save_base_url,
)


app = Flask(__name__)
app.secret_key = os.environ.get("FLASK_SECRET_KEY", "ppt-secret-key")
OUTPUT_ROOT = get_output_root()
STATE_SESSION_KEY = "ppt_state"
SESSION_ID_KEY = "ppt_conversation_id"


@lru_cache(maxsize=8)
def _cached_store(root):
    return PPTStore(root)


def _store():
    return _cached_store(str(OUTPUT_ROOT.resolve()))


def _owner():
    # 本地单账号占位；云端上线前需接入经过确认的登录与授权方案。
    return LOCAL_OWNER


def _create_run_dir() -> Path:
    run_id = datetime.now().strftime("%Y%m%d_%H%M%S")
    OUTPUT_ROOT.mkdir(parents=True, exist_ok=True)
    return Path(mkdtemp(prefix=f"{run_id}_", dir=OUTPUT_ROOT))


def _parse_page_numbers(raw: str, max_page: Optional[int] = None):
    raw = (raw or "").strip()
    if not raw or len(raw) > 2000:
        return None
    raw = re.sub(r"(?<=\d)\s*[-–—－]\s*(?=\d)", "-", raw)
    pages = set()
    for part in re.split(r"[,\s，;；]+", raw):
        if re.fullmatch(r"\d+", part):
            number = int(part)
            if number < 1 or (max_page is not None and number > max_page):
                return None
            pages.add(number)
        elif re.fullmatch(r"\d+-\d+", part):
            start, end = sorted(int(x) for x in part.split("-"))
            if start < 1 or (max_page is not None and end > max_page) or end - start > 10000:
                return None
            pages.update(range(start, end + 1))
        elif part:
            return None
    return sorted(pages) or None


def _safe_ppt_filename(filename: str) -> str:
    leaf = re.split(r"[\\/]", filename)[-1]
    stem = re.sub(r"\.pptx$", "", leaf, flags=re.IGNORECASE)
    stem = re.sub(r"[^\w-]+", "_", stem).strip("_")[:80] or "generated_presentation"
    if re.fullmatch(r"CON|PRN|AUX|NUL|COM[1-9]|LPT[1-9]", stem, flags=re.IGNORECASE):
        stem = f"presentation_{stem}"
    return f"{stem}.pptx"


def _incremental_label(use_incremental, pages, previous, scope):
    text = {"text": "仅改文本", "design": "仅改设计", "both": "整体修改"}[scope]
    total = len((previous.get("outline") or {}).get("slides") or [])
    if use_incremental and len(pages or []) == total:
        return f"全局{text}（全部 {total} 页）"
    return f"{text}（第 {', '.join(map(str, pages or []))} 页）"


def _load_state():
    # 旧 Cookie 的整份状态不再作为数据来源；归档通过显式导入读取。
    session.pop(STATE_SESSION_KEY, None)
    conversation_id = session.get(SESSION_ID_KEY)
    if not conversation_id:
        return PPTGenerationState(topic="")
    return _store().load_state(conversation_id, _owner())


def _page(state, **extra):
    conversation = _store().conversation(state.get("_conversation_id"), _owner()) if state.get("_conversation_id") else None
    values = extra.pop("form_values", None) or {"topic": conversation["title"] if conversation else state.topic,
        "requirements": state.requirements, "filename": "generated_presentation", "page_range": "",
        "modify_page": ",".join(map(str, state.modify_pages)), "mode": state.get("requested_mode", state.operation_mode)}
    return render_template(
        "index.html", state=state, current_outline=state.get("outline"),
        has_previous_state=bool(state.get("outline") or state.get("blueprint")),
        request_id=str(uuid4()), conversations=_store().conversations(_owner()), form_values=values, **extra,
    )


def _build_ppt_preview(ppt_path, outline=None):
    try:
        presentation = Presentation(ppt_path)
    except Exception:
        return []
    preview = []
    titles = (outline or {}).get("slides", [])
    for i, slide in enumerate(presentation.slides):
        texts = []
        for shape in slide.shapes:
            if shape.has_table:
                texts.extend(" / ".join(cell.text for cell in row.cells) for row in shape.table.rows)
            elif hasattr(shape, "text") and shape.text:
                texts.append(shape.text.strip())
        title = titles[i].get("title") if i < len(titles) and isinstance(titles[i], dict) else None
        title = title or (texts[0] if texts else "未命名页")
        if title in texts:
            texts.remove(title)
        preview.append({"title": title, "text": " | ".join(texts)})
    return preview


@app.after_request
def prevent_state_cache(response):
    if request.path == "/" or request.path in {"/state", "/generate"} or request.path.startswith("/api/"):
        response.headers["Cache-Control"] = "private, no-store"
    return response


@app.get("/")
def index():
    conversation_id = request.args.get("conversation_id")
    if conversation_id:
        if not _store().conversation(conversation_id, _owner()):
            return _page(PPTGenerationState(topic=""), error="文稿不存在"), 404
        return _page(_store().load_state(conversation_id, _owner()))
    return _page(_load_state())


@app.get("/setup")
def setup_page():
    return render_template("setup.html", has_key=bool(os.environ.get("DASHSCOPE_API_KEY") or has_api_key()))


@app.get("/api/has_key")
def api_has_key():
    return jsonify({"has_key": bool(os.environ.get("DASHSCOPE_API_KEY") or has_api_key()), "base_url": load_base_url()})


@app.post("/api/save_key")
def api_save_key():
    payload = request.get_json(silent=True) or {}
    key = (payload.get("api_key") or request.form.get("api_key") or "").strip()
    base_url = (payload.get("base_url") or request.form.get("base_url") or "").strip()
    if not key:
        return jsonify({"status": "error", "message": "请输入 API Key"}), 400
    try:
        save_api_key(key)
    except Exception as exc:
        return jsonify({"status": "error", "message": f"密钥加密保存失败：{type(exc).__name__}"}), 500
    save_base_url(base_url)
    os.environ["DASHSCOPE_API_KEY"] = key
    if base_url:
        os.environ["DASHSCOPE_BASE_URL"] = base_url
    else:
        os.environ.pop("DASHSCOPE_BASE_URL", None)
    return jsonify({"status": "success"})


@app.post("/api/logout")
def api_logout():
    clear_api_key()
    os.environ.pop("DASHSCOPE_API_KEY", None)
    return jsonify({"status": "success"})


@app.post("/api/reset")
def reset_conversation():
    session.pop(SESSION_ID_KEY, None)
    session.pop(STATE_SESSION_KEY, None)
    return jsonify({"status": "success", "message": "已打开空白会话，历史仍可从列表选择。"})


@app.get("/api/conversations")
def conversation_list():
    return jsonify({"conversations": _store().conversations(_owner())})


@app.post("/api/conversations/<conversation_id>/select")
def select_conversation(conversation_id):
    if not _store().conversation(conversation_id, _owner()):
        return jsonify({"status": "error", "message": "会话不存在"}), 404
    session.pop(STATE_SESSION_KEY, None)
    session[SESSION_ID_KEY] = conversation_id
    return jsonify({"status": "success"})


@app.get("/api/conversations/<conversation_id>/versions")
def version_list(conversation_id):
    store = _store()
    if not store.conversation(conversation_id, _owner()):
        return jsonify({"message": "会话不存在"}), 404
    versions = store.versions(conversation_id, _owner())
    for version in versions:
        version["artifacts"] = [{"id": item["id"], "role": item["role"], "sha256": item["sha256"],
                                 "download_url": f"/api/artifacts/{item['id']}/download"}
                                for item in store.version_artifacts(version["id"], _owner())]
    return jsonify({"versions": versions})


@app.get("/api/artifacts/<artifact_id>/download")
def download_artifact(artifact_id):
    store = _store()
    artifact = store.artifact(artifact_id, _owner())
    if not artifact:
        return jsonify({"message": "文件不存在"}), 404
    if not store.verify_assets([artifact]):
        return jsonify({"message": "文件缺失或校验失败，请从备份恢复"}), 409
    path = store.asset_path(artifact["path"])
    inline = request.args.get("inline") == "1" and artifact["role"] in {"preview", "pdf"}
    response = send_file(path, as_attachment=not inline, download_name=path.name, max_age=0)
    response.headers["X-Content-Type-Options"] = "nosniff"
    return response


@app.get("/api/jobs/<job_id>")
def job_status(job_id):
    store = _store()
    job = store.job(job_id, _owner())
    if not job:
        return jsonify({"message": "任务不存在"}), 404
    fields = ("id", "conversation_id", "base_version_id", "result_version_id", "status", "error", "created_at", "updated_at")
    result = {field: job[field] for field in fields}
    result["stages"] = [{field: stage[field] for field in
                         ("stage_key", "input_hash", "prompt_version", "model", "status", "request_id", "created_at", "updated_at")}
                        | {"calls": json.loads(stage["usage_json"] or "[]")} for stage in store.stages(job_id)]
    return jsonify(result)


@app.post("/api/jobs/<job_id>/retry")
def retry_job(job_id):
    if not _store().job(job_id, _owner()):
        return jsonify({"message": "任务不存在"}), 404
    payload = request.get_json(silent=True) or {}
    try:
        _store().retry_job(job_id, confirm_unknown=payload.get("confirm_unknown") is True, owner=_owner())
    except StorageConflict as exc:
        return jsonify({"message": str(exc)}), 409
    try:
        dispatch_pending(OUTPUT_ROOT)
    except Exception:
        app.logger.warning("Task enqueue delayed; worker will recover pending tasks")
    return jsonify({"status": "queued", "job_id": job_id}), 202


@app.get("/api/archives")
def archive_list():
    from config.archives import scan_archives
    return jsonify({"archives": scan_archives(_store(), _owner())})


@app.post("/api/archives/<run_id>/import")
def import_archive(run_id):
    from config.archives import read_archive
    try:
        snapshot, manifest, source_hash = read_archive(_store(), run_id)
        expected = (request.get_json(silent=True) or {}).get("source_hash")
        if expected and expected != source_hash:
            return jsonify({"message": "归档内容已改变，请重新扫描后导入"}), 409
        conversation_id = _store().import_legacy(run_id, source_hash, snapshot, manifest, _owner())
    except (ValueError, OSError) as exc:
        return jsonify({"message": str(exc) if isinstance(exc, ValueError) else "归档无法读取"}), 400
    session[SESSION_ID_KEY] = conversation_id
    return jsonify({"status": "success", "conversation_id": conversation_id})


def _design_public(job):
    store = _store()
    data = json.loads(job["snapshot_json"]).get("data", {}).get("design_candidates") or {}
    candidates = []
    for item in data.get("candidates", []):
        assets = [{key: asset.get(key) for key in ("id", "role", "page", "sha256")} | {
            "download_url": f"/api/design/jobs/{job['id']}/assets/{asset['id']}",
            "image_url": f"/api/design/jobs/{job['id']}/assets/{asset['id']}?inline=1" if asset["role"] == "preview" else ""}
            for asset in item.get("assets", [])]
        valid = item.get("status") == "ready" and store.verify_assets(item.get("assets", []))
        candidates.append({key: item.get(key) for key in ("id", "name", "description", "status", "reason", "rules", "icon_count")} | {"assets": assets, "integrity": valid})
    baseline_assets = [{key: asset.get(key) for key in ("id", "page", "sha256")} | {"image_url": f"/api/design/jobs/{job['id']}/assets/{asset['id']}?inline=1"}
                       for asset in data.get("baseline_assets", [])]
    return {key: job[key] for key in ("id", "status", "base_version_id", "error", "created_at")} | {
        "pages": data.get("pages", json.loads(job["request_json"]).get("page_numbers", [])), "candidates": candidates, "baseline_assets": baseline_assets,
        "calls_used": data.get("calls_used", 0), "max_layout_calls": data.get("max_layout_calls", 12)}


@app.get("/api/conversations/<conversation_id>/design")
def design_list(conversation_id):
    store = _store()
    if not store.conversation(conversation_id, _owner()):
        return jsonify({"message": "文稿不存在"}), 404
    return jsonify({"jobs": [_design_public(job) for job in store.design_jobs(conversation_id, _owner())]})


@app.get("/api/design/jobs/<job_id>/assets/<asset_id>")
def design_asset(job_id, asset_id):
    store = _store(); job = store.job(job_id, _owner())
    if not job or json.loads(job["request_json"]).get("kind") != "design_candidates":
        return jsonify({"message": "候选不存在"}), 404
    data = json.loads(job["snapshot_json"]).get("data", {}).get("design_candidates") or {}
    all_assets = data.get("baseline_assets", []) + [asset for candidate in data.get("candidates", []) for asset in candidate.get("assets", [])]
    asset = next((asset for asset in all_assets if asset["id"] == asset_id), None)
    if not asset:
        return jsonify({"message": "候选文件不存在"}), 404
    if not store.verify_assets([asset]):
        return jsonify({"message": "候选文件缺失或校验失败"}), 409
    path = store.asset_path(asset["path"])
    inline = request.args.get("inline") == "1" and asset["role"] in {"preview", "pdf"}
    response = send_file(path, as_attachment=not inline, download_name=path.name, max_age=0)
    response.headers["X-Content-Type-Options"] = "nosniff"
    return response


def _submit_design(payload, kind):
    store = _store()
    conversation_id = payload.get("conversation_id"); version_id = payload.get("version_id"); key = payload.get("request_id")
    if not isinstance(conversation_id, str) or not isinstance(version_id, str) or not isinstance(key, str) or not re.fullmatch(r"[A-Za-z0-9_-]{8,128}", key):
        return jsonify({"message": "文稿、版本或请求标识无效"}), 400
    conversation = store.conversation(conversation_id, _owner()); version = store.version(version_id, _owner())
    if not conversation or not version or version["conversation_id"] != conversation_id:
        return jsonify({"message": "文稿或版本不存在"}), 404
    state = PPTGenerationState.from_snapshot(store.rebase_snapshot(json.loads(version["snapshot_json"])))
    if not isinstance(state.get("outline"), dict) or not isinstance(state.get("blueprint"), dict):
        return jsonify({"message": "旧文稿缺少语义大纲或物理蓝图，请先重新生成"}), 400
    request_data = {"kind": kind}
    if kind == "design_candidates":
        pages = payload.get("pages"); design_request = payload.get("design_request", "")
        count = len(state.get("outline").get("slides") or [])
        if not isinstance(pages, list) or not 1 <= len(pages) <= 3 or len(set(map(str, pages))) != len(pages) or any(type(n) is not int or n < 1 or n > count for n in pages):
            return jsonify({"message": "请选择 1–3 个有效页面"}), 400
        if not isinstance(design_request, str) or len(design_request) > 2000:
            return jsonify({"message": "设计要求最多 2000 字"}), 400
        request_data.update(page_numbers=sorted(pages), design_request=design_request.strip(), user_input="生成设计候选：" + design_request.strip())
    else:
        source_id = payload.get("source_job_id"); candidate_id = payload.get("candidate_id")
        source = store.job(source_id, _owner()) if isinstance(source_id, str) else None
        if not source or source["conversation_id"] != conversation_id or source["base_version_id"] != version_id or json.loads(source["request_json"]).get("kind") != "design_candidates":
            return jsonify({"message": "候选不属于所选文稿版本"}), 400
        data = json.loads(source["snapshot_json"]).get("data", {}).get("design_candidates") or {}
        candidate = next((item for item in data.get("candidates", []) if item["id"] == candidate_id), None)
        if not candidate or candidate["status"] != "ready" or not store.verify_assets(candidate["assets"]):
            return jsonify({"message": "候选尚未就绪或文件校验失败"}), 409
        candidate_hash = next(asset["sha256"] for asset in candidate["assets"] if asset["role"] == "blueprint")
        if payload.get("candidate_hash") != candidate_hash:
            return jsonify({"message": "候选已更新，请重新预览后采用"}), 409
        request_data.update(source_job_id=source_id, candidate_id=candidate_id, candidate_hash=candidate_hash,
                            page_numbers=data["pages"], user_input="采用设计候选：" + candidate["name"])
    fingerprint = content_hash({"conversation_id": conversation_id, "version_id": version_id, **request_data})
    existing = store.job_by_key(key, _owner())
    if existing:
        if existing["request_hash"] != fingerprint:
            return jsonify({"message": "同一请求标识不能用于不同设计操作"}), 409
        return jsonify({"job_id": existing["id"], "conversation_id": conversation_id, "status": existing["status"]}), 202
    if conversation["current_version_id"] != version_id:
        return jsonify({"message": "当前版本已改变，请重新打开文稿后操作"}), 409
    state.set("design_candidates", None)
    run_dir = _create_run_dir()
    try:
        job, created = store.create_job(request_data, state.snapshot(), run_dir, key, fingerprint, conversation_id, version_id, _owner())
    except StorageConflict as exc:
        run_dir.rmdir(); return jsonify({"message": str(exc)}), 409
    if not created:
        run_dir.rmdir()
    try:
        dispatch_pending(OUTPUT_ROOT)
    except Exception:
        app.logger.warning("Design enqueue delayed; worker will recover pending tasks")
    return jsonify({"job_id": job["id"], "conversation_id": conversation_id, "status": job["status"]}), 202


@app.post("/api/design/candidates")
def submit_design_candidates():
    payload = request.get_json(silent=True)
    if not isinstance(payload, dict):
        return jsonify({"message": "设计请求格式无效"}), 400
    return _submit_design(payload, "design_candidates")


@app.post("/api/design/adopt")
def adopt_design_candidate():
    payload = request.get_json(silent=True)
    if not isinstance(payload, dict):
        return jsonify({"message": "设计请求格式无效"}), 400
    return _submit_design(payload, "design_adopt")


def _version_view(store, version_id):
    version = store.version(version_id, _owner()) if version_id else None
    if not version:
        return {"preview": [], "visual_review": {"status": "not_run", "pages": []},
                "download_url": "", "last_file": "", "artifact_integrity": "none", "layout_fallbacks": []}
    artifacts = store.version_artifacts(version_id, _owner())
    ppt = next((item for item in artifacts if item["role"] == "pptx"), None)
    last_file = str(store.asset_path(ppt["path"])) if ppt else ""
    valid = bool(ppt and store.verify_assets([ppt]))
    committed = json.loads(version["snapshot_json"]).get("data", {})
    preview = _build_ppt_preview(last_file, committed.get("outline")) if valid else []
    images = {}
    for item in artifacts:
        match = re.fullmatch(r"page-(\d+)", Path(item["path"]).stem)
        if item["role"] == "preview" and match:
            images[int(match[1])] = item
    visual = committed.get("visual_review") or {"status": "not_run", "pages": []}
    for number, page in enumerate(preview, 1):
        image = images.get(number)
        if image and store.verify_assets([image]):
            page["image_url"] = f"/api/artifacts/{image['id']}/download?inline=1"
        review = next((item for item in visual.get("pages", []) if item.get("page") == number), {})
        page.update(review_status=review.get("status", "not_run"), issues=review.get("issues", []), rule_issues=review.get("rule_issues", []))
    return {"preview": preview, "visual_review": visual, "last_file": last_file,
            "layout_fallbacks": committed.get("layout_fallbacks", []),
            "artifact_integrity": "valid" if valid else ("invalid" if ppt else "none"),
            "download_url": f"/api/artifacts/{ppt['id']}/download" if valid else ""}


@app.get("/api/versions/<version_id>/preview")
def version_preview(version_id):
    version = _store().version(version_id, _owner())
    if not version:
        return jsonify({"message": "版本不存在"}), 404
    return jsonify({**_version_view(_store(), version_id), "version_id": version_id, "conversation_id": version["conversation_id"]})


@app.get("/state")
def state_api():
    store = _store(); requested = request.args.get("conversation_id")
    if requested and not store.conversation(requested, _owner()):
        return jsonify({"message": "会话不存在"}), 404
    state = store.load_state(requested, _owner()) if requested else _load_state()
    conversation = store.conversation(state.get("_conversation_id"), _owner()) if state.get("_conversation_id") else None
    return jsonify({**_version_view(store, state.get("_version_id")),
        "status": state.status, "generation_status": state.get("generation_status", "idle"),
        "result_kind": state.get("result_kind", ""), "content_review_status": state.get("content_review_status", "not_run"),
        "layout_recovery_request": state.get("layout_recovery_request"),
        "progress_log": state.progress_log, "last_updated": state.last_updated,
        "conversation_id": state.get("_conversation_id"), "version_id": state.get("_version_id"),
        "job_id": state.get("_job_id"), "topic": state.topic,
        "conversation_title": conversation["title"] if conversation else "新文稿",
    })


@app.get("/api/conversations/<conversation_id>/dialogue")
def dialogue_history(conversation_id):
    store = _store()
    if not store.conversation(conversation_id, _owner()):
        return jsonify({"message": "会话不存在"}), 404
    cursor = request.args.get("before_id")
    if cursor and (not cursor.isdigit() or len(cursor) > 18 or int(cursor) < 1):
        return jsonify({"message": "历史记录游标无效"}), 400
    records = store.dialogue_records(conversation_id, _owner(), limit=101, before_id=int(cursor) if cursor else None)
    visible = [{key: item.get(key) for key in ("id", "role", "content", "kind", "job_id", "response", "version_id", "created_at")}
               for item in records[-100:]]
    return jsonify({"messages": visible, "jobs": store.dialogue_jobs(conversation_id, _owner()), "has_older": len(records) > 100, "before_id": visible[0]["id"] if visible else None})


@app.post("/api/dialogue")
def submit_dialogue():
    payload = request.get_json(silent=True)
    if not isinstance(payload, dict):
        return jsonify({"message": "请发送有效的对话请求"}), 400
    message = payload.get("message"); intent = payload.get("intent", "discuss")
    if not isinstance(message, str) or not message.strip() or len(message) > 8000 or not isinstance(intent, str) or intent not in {"discuss", "question"}:
        return jsonify({"message": "请输入 1–8000 字的消息，并选择讨论需求或提问内容"}), 400
    declared = payload.get("conversation_id") or None
    if declared is not None and not isinstance(declared, str):
        return jsonify({"message": "会话标识无效"}), 400
    conversation_id = declared or session.get(SESSION_ID_KEY)
    store = _store(); conversation = store.conversation(conversation_id, _owner()) if conversation_id else None
    if conversation_id and not conversation:
        return jsonify({"message": "会话不存在"}), 404
    version_id = payload.get("version_id") or (conversation["current_version_id"] if conversation else None)
    version = store.version(version_id, _owner()) if isinstance(version_id, str) else None
    if version_id and (not version or version["conversation_id"] != conversation_id):
        return jsonify({"message": "所选版本不属于当前会话"}), 400
    if intent == "question" and not version:
        return jsonify({"message": "请先生成或选择一份已保存的 PPT，再提问内容"}), 400
    pages = payload.get("pages") or []
    count = len((json.loads(version["snapshot_json"]).get("data", {}).get("outline") or {}).get("slides") or []) if version else 0
    if not isinstance(pages, list) or len(pages) > 50 or any(type(n) is not int or n < 1 or n > count for n in pages):
        return jsonify({"message": "提问页面范围无效"}), 400
    key = payload.get("request_id")
    if not isinstance(key, str) or not re.fullmatch(r"[A-Za-z0-9_-]{8,128}", key):
        return jsonify({"message": "对话请求标识无效"}), 400
    fingerprint = content_hash({"message": message.strip(), "intent": intent, "conversation_id": declared,
                                "version_id": payload.get("version_id"), "pages": pages})
    existing = store.job_by_key(key, _owner())
    if existing:
        if existing["request_hash"] != fingerprint:
            return jsonify({"message": "同一请求标识不能用于不同消息"}), 409
        session[SESSION_ID_KEY] = existing["conversation_id"]
        return jsonify({"job_id": existing["id"], "conversation_id": existing["conversation_id"], "status": existing["status"]}), 202
    state = PPTGenerationState(topic=conversation["title"] if conversation else message.strip()[:80])
    through_id = store.latest_message_id(conversation_id, _owner()) if conversation_id else 0
    request_data = {"kind": "dialogue", "intent": intent, "user_input": message.strip(), "page_numbers": sorted(set(pages)),
                    "requirements": "", "history_through_id": through_id}
    run_dir = _create_run_dir()
    try:
        job, created = store.create_job(request_data, state.snapshot(), run_dir, key, fingerprint, conversation_id, version_id, _owner())
    except StorageConflict as exc:
        run_dir.rmdir()
        return jsonify({"message": str(exc)}), 409
    if not created:
        run_dir.rmdir()
    session[SESSION_ID_KEY] = job["conversation_id"]
    try:
        dispatch_pending(OUTPUT_ROOT)
    except Exception:
        app.logger.warning("Dialogue enqueue delayed; worker will recover pending tasks")
    return jsonify({"job_id": job["id"], "conversation_id": job["conversation_id"], "status": job["status"]}), 202


@app.post("/generate")
def generate():
    store = _store()
    previous_state = _load_state()
    topic = (request.form.get("topic") or "").strip()
    requirements = (request.form.get("requirements") or "").strip()
    safe_name = _safe_ppt_filename((request.form.get("filename") or "generated_presentation.pptx").strip())
    mode = (request.form.get("mode") or "new").strip()
    raw_pages = (request.form.get("modify_page") or "").strip()
    raw_range = (request.form.get("page_range") or "").strip()

    def invalid(message, status=400):
        return _page(previous_state, error=message, form_values={key: request.form.get(key, "") for key in ("topic", "requirements", "filename", "page_range", "modify_page", "mode")}), status

    if not topic:
        return invalid("请输入主题内容。")
    if mode not in {"new", "continue", "modify_text", "modify_design"}:
        return invalid("不支持的生成模式。")
    try:
        policy = parse_page_policy(raw_range, topic + "\n" + requirements) if mode == "new" else (parse_page_policy(raw_range) if raw_range else None)
    except ContentConstraintError as exc:
        return invalid(str(exc))
    conversation_id = request.form.get("conversation_id") or session.get(SESSION_ID_KEY)
    if mode == "new":
        draft = store.conversation(conversation_id, _owner()) if conversation_id else None
        conversation_id = conversation_id if draft and not draft["current_version_id"] and not store.latest_job(conversation_id, _owner()) else None
    base_version = None if mode == "new" else request.form.get("base_version_id", previous_state.get("_version_id"))
    key = request.headers.get("Idempotency-Key") or request.form.get("request_id") or str(uuid4())
    if not re.fullmatch(r"[A-Za-z0-9_-]{8,128}", key):
        return invalid("请求标识格式无效，请刷新后提交。")
    fingerprint = content_hash({"topic": topic, "requirements": requirements, "filename": safe_name,
                                "mode": mode, "pages": raw_pages, "page_range": raw_range,
                                "conversation_id": (request.form.get("conversation_id") or None) if mode == "new" else conversation_id, "base_version": base_version})
    existing = store.job_by_key(key, _owner())
    if existing:
        if existing["request_hash"] != fingerprint:
            return invalid("同一请求标识不能用于不同输入。", 409)
        session[SESSION_ID_KEY] = existing["conversation_id"]
        return _page(_load_state(), success=True, file_name=safe_name, mode_label="已提交任务")

    if mode != "new":
        conversation = store.conversation(conversation_id, _owner())
        if not conversation or not conversation["current_version_id"]:
            return invalid("当前没有可修改的已提交版本，请先生成一份 PPT。")
        if base_version != conversation["current_version_id"]:
            return invalid("当前版本已改变，请刷新后重新提交。", 409)
        version = store.version(base_version, _owner())
        previous_state = PPTGenerationState.from_snapshot(store.rebase_snapshot(json.loads(version["snapshot_json"])))
    outline = previous_state.get("outline")
    slides = outline.get("slides") if isinstance(outline, dict) else None
    if mode != "new" and not (isinstance(slides, list) and slides):
        return invalid("当前没有可修改的大纲缓存，请先生成一份 PPT。")
    pages = _parse_page_numbers(raw_pages, len(slides)) if mode != "new" and raw_pages else None
    if mode != "new" and raw_pages and pages is None:
        return invalid(f"页码无效，请输入 1–{len(slides)} 内的页码或范围。")
    incremental = mode != "new" and (pages is not None or mode in {"modify_text", "modify_design"})
    if incremental:
        blueprint = previous_state.get("blueprint")
        canvases = blueprint.get("canvases") if isinstance(blueprint, dict) else None
        if not isinstance(canvases, list) or len(canvases) != len(slides):
            return invalid("当前物理蓝图缓存缺失或页数不一致，无法安全执行页面修改。")
        pages = pages or list(range(1, len(slides) + 1))
        if policy:
            try:
                check_page_count(len(slides), policy)
            except ContentConstraintError:
                return invalid("页面修改不能改变总页数；请选择重新生成来调整页数。")
    if mode != "new" and policy is None:
        policy = {"min_slides": len(slides), "max_slides": len(slides)}
    else:
        policy = policy.model_dump()
    user_input = f"{topic}\n要求：{requirements}" if requirements else topic
    state = PPTGenerationState(topic=topic, requirements=requirements) if mode == "new" else previous_state
    state.topic, state.requirements = topic, requirements
    scope = {"modify_text": "text", "modify_design": "design"}.get(mode, "both")
    state.set_operation("incremental_page" if incremental else mode, scope=scope,
                        page=pages[0] if pages else None, pages=pages)
    state.add_requirement_history(requirements)
    state.append_message("user", user_input)
    state.set("generation_status", "queued")
    state.set("result_kind", "")
    state.set("requested_mode", mode)
    state.set("page_policy", policy)
    state.append_progress("Job: 已保存任务，等待 worker 执行")
    through_id = store.latest_message_id(conversation_id, _owner()) if conversation_id else 0
    payload = {"user_input": user_input, "requirements": requirements, "filename": safe_name,
               "use_incremental": incremental, "page_numbers": pages, "mode": mode,
               "history_through_id": through_id}
    run_dir = _create_run_dir()
    try:
        job, created = store.create_job(payload, state.snapshot(), run_dir, key, fingerprint,
                                        conversation_id, base_version, _owner())
    except StorageConflict as exc:
        run_dir.rmdir()
        return invalid(str(exc), 409)
    if not created:
        run_dir.rmdir()
    session[SESSION_ID_KEY] = job["conversation_id"]
    try:
        dispatch_pending(OUTPUT_ROOT)
    except Exception:
        app.logger.warning("Task enqueue delayed; worker will recover pending tasks")
    return _page(_load_state(), success=True, file_name=safe_name,
                 mode_label=_incremental_label(incremental, pages, previous_state, scope) if incremental else
                 ("重新生成" if mode == "new" else "基于历史继续修改"))


if __name__ == "__main__":
    app.run(host="127.0.0.1", port=8000, debug=False)
