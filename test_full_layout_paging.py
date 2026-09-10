# -*- coding: utf-8 -*-
"""
全量管线「逐页排版」验证（不依赖真实 LLM API）。

背景：全量生成时 Layout Agent 一次性生成全部页物理蓝图，单次请求输出过大/耗时过长，
触发 MaaS 网关主动断连（Connection error）。修复后改为逐页排版：每次仅提交单页语义大纲。

本测试通过 from_content_json 注入 3 页大纲（跳过 Planner/Writer/生图），仅验证 Layout 逐页逻辑：
1. Layout 被逐页调用 N 次（每次输入仅 1 页）；
2. 最终蓝图拼接为 N 页 canvas；
3. state["blueprint"] 被显式回写为全量蓝图；
4. 渲染出的 pptx 页数正确。
"""
import json
import os
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
os.environ["DASHSCOPE_API_KEY"] = "fake-key-for-local-test"

from pptx import Presentation

from config.schema import (
    CanvasElement,
    PPTFinalBlueprint,
    PPTOutline,
    SlideCanvas,
    TextParagraph,
    ColorPalette,
)
from config.state import PPTGenerationState
import main as main_module

TEST_OUT_DIR = Path(__file__).resolve().parent / "debug_logs" / "full_layout_test"
TEST_OUT_DIR.mkdir(parents=True, exist_ok=True)

THEME = ColorPalette(primary="#1E3A8A", secondary="#10B981", content="#64748B", background="#F8FAFC")


def make_slide(title, idx):
    return {
        "title": title,
        "punchline": f"页面{idx}摘要",
        "bullet_points": [f"要点{idx}A", f"要点{idx}B"],
        "full_content": [f"正文{idx}A", f"正文{idx}B"],
        "theme": THEME.model_dump(),
        "layout_type": "standard_text",
        "image_prompt": None,
        "image_path": None,
        "icons": [],
        "svg_code": None,
    }


LAYOUT_CALLS = {"count": 0, "received_titles": []}


def install_fake_layout():
    def fake_get_layout_agent(api_key, model_name="x"):
        def layout(state, writer_output_json):
            LAYOUT_CALLS["count"] += 1
            slides = writer_output_json.get("slides") or []
            # 关键断言：每次输入必须是单页
            assert len(slides) == 1, f"Layout 第 {LAYOUT_CALLS['count']} 次调用应只收到单页，实际 {len(slides)} 页"
            slide = slides[0]
            LAYOUT_CALLS["received_titles"].append(slide["title"])
            elements = [
                CanvasElement(
                    type="textbox", left=0.8, top=0.8, width=11.5, height=1.2,
                    paragraphs=[TextParagraph(text=slide["title"], font_size=24, bold=True)],
                ),
                CanvasElement(
                    type="textbox", left=0.8, top=2.2, width=11.5, height=4.0,
                    paragraphs=[TextParagraph(text=" | ".join(slide.get("full_content") or []), font_size=14)],
                ),
            ]
            return PPTFinalBlueprint(canvases=[SlideCanvas(background_color="#FFFFFF", elements=elements)])

        return layout

    main_module.get_layout_agent = fake_get_layout_agent


def run():
    print("=" * 60)
    print("TEST: 全量管线逐页排版（注入 3 页大纲，Layout 应逐页调用 3 次）")
    print("=" * 60)

    install_fake_layout()

    # 构造 3 页大纲并落盘为 from_content_json
    slides = [make_slide(f"第{i}页标题", i) for i in range(1, 4)]
    outline_dict = {"topic": "测试主题", "slides": slides}
    outline_json = TEST_OUT_DIR / "outline.json"
    outline_json.write_text(json.dumps(outline_dict, ensure_ascii=False), encoding="utf-8")

    state = PPTGenerationState(topic="测试主题", requirements="")
    output = str(TEST_OUT_DIR / "full_layout_paging.pptx")

    result = main_module.run_ppt_system(
        "测试主题",
        filename=output,
        from_content_json=str(outline_json),
        state=state,
        requirements="",
        run_dir=str(TEST_OUT_DIR / "run"),
    )

    assert result == output, f"应返回输出路径，实际: {result}"
    assert LAYOUT_CALLS["count"] == 3, f"Layout 应被逐页调用 3 次，实际 {LAYOUT_CALLS['count']}"
    assert LAYOUT_CALLS["received_titles"] == ["第1页标题", "第2页标题", "第3页标题"], \
        f"逐页顺序应为 1→2→3，实际 {LAYOUT_CALLS['received_titles']}"

    # 最终蓝图应拼接为 3 页
    blueprint = state.get("blueprint")
    assert blueprint and len(blueprint.get("canvases") or []) == 3, \
        f"state.blueprint 应回写为 3 页画布，实际 {len((blueprint or {}).get('canvases') or [])}"

    prs = Presentation(output)
    assert len(prs.slides) == 3, f"渲染出的 pptx 应为 3 页，实际 {len(prs.slides)}"

    print("PASS: 全量管线逐页排版——Layout 逐页调用 3 次、蓝图拼接 3 页、state 回写、pptx 3 页\n")


if __name__ == "__main__":
    run()
    print("ALL FULL-LAYOUT PAGING TESTS PASSED")
